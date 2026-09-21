#!/usr/bin/env python3
"""
CU-A1 验收任务：生成 sales-audit.xlsx
运行方式：python generate_xlsx.py
依赖：pip install openpyxl
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers

DIR = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(DIR, "sales-audit.xlsx")

# ── 原始数据（B修订后 + C退货） ──────────────────────────
# 订单, 地区, 商品, 最终数量, 单价, 退货数量, 状态
ROWS = [
    ("O101", "华东", "键盘", 3, 120, 1, "正常"),
    ("O102", "华南", "鼠标", 4,  80, 0, "正常"),
    ("O103", "华东", "支架", 2, 150, 0, "正常"),
    ("O104", "华北", "键盘", 0,   0, 0, "取消"),
    ("O105", "华南", "支架", 1, 150, 0, "正常"),
    ("O106", "华北", "鼠标", 6,  80, 2, "正常"),
    ("O107", "华东", "鼠标", 2,  80, 0, "正常"),
    ("O108", "华南", "键盘", 2, 110, 0, "正常"),
    ("O109", "华北", "支架", 3, 150, 0, "正常"),
]

wb = Workbook()

# ── 样式 ─────────────────────────────────────────────────
header_font = Font(bold=True)
header_fill = PatternFill("solid", fgColor="4472C4")
header_font_w = Font(bold=True, color="FFFFFF")
thin_border = Border(
    left=Side("thin"), right=Side("thin"),
    top=Side("thin"), bottom=Side("thin"),
)
money_fmt = '#,##0'

def style_header(ws, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font_w
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

def style_data(ws, rows, cols):
    for r in range(2, rows + 2):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = thin_border
            if c in (5, 8, 9, 10):  # 金额列
                cell.number_format = money_fmt
            cell.alignment = Alignment(horizontal="center")

# ══════════════════════════════════════════════════════════
# Sheet 1: 订单明细
# ══════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "订单明细"

headers1 = ["订单号", "地区", "商品", "最终数量", "单价", "退货数量", "状态", "毛额", "退款", "净额"]
ws1.append(headers1)
style_header(ws1, len(headers1))

for i, (oid, region, prod, qty, price, ret, status) in enumerate(ROWS, start=2):
    ws1.cell(row=i, column=1, value=oid)
    ws1.cell(row=i, column=2, value=region)
    ws1.cell(row=i, column=3, value=prod)
    ws1.cell(row=i, column=4, value=qty)
    ws1.cell(row=i, column=5, value=price)
    ws1.cell(row=i, column=6, value=ret)
    ws1.cell(row=i, column=7, value=status)
    # 毛额 = 最终数量 × 单价（公式）
    ws1.cell(row=i, column=8).value = f"=D{i}*E{i}"
    # 退款 = 退货数量 × 单价（公式）
    ws1.cell(row=i, column=9).value = f"=F{i}*E{i}"
    # 净额 = 毛额 - 退款（公式）
    ws1.cell(row=i, column=10).value = f"=H{i}-I{i}"

# 合计行
last = len(ROWS) + 1
tr = last + 1
ws1.cell(row=tr, column=1, value="合计")
ws1.cell(row=tr, column=1).font = Font(bold=True)
ws1.cell(row=tr, column=8).value = f"=SUM(H2:H{last})"
ws1.cell(row=tr, column=9).value = f"=SUM(I2:I{last})"
ws1.cell(row=tr, column=10).value = f"=SUM(J2:J{last})"
for c in (8, 9, 10):
    ws1.cell(row=tr, column=c).font = Font(bold=True)
    ws1.cell(row=tr, column=c).number_format = money_fmt

style_data(ws1, len(ROWS), len(headers1))
# 列宽
for col, w in zip("ABCDEFGHIJ", [8, 6, 6, 8, 8, 8, 6, 10, 10, 10]):
    ws1.column_dimensions[col].width = w

# ══════════════════════════════════════════════════════════
# Sheet 2: 区域汇总（按净额降序）
# ══════════════════════════════════════════════════════════
ws2 = wb.create_sheet("区域汇总")
headers2 = ["地区", "订单数", "毛额", "退款", "净额"]
ws2.append(headers2)
style_header(ws2, len(headers2))

# 区域按净额降序：华北、华东、华南
regions_order = ["华北", "华东", "华南"]
for i, reg in enumerate(regions_order, start=2):
    ws2.cell(row=i, column=1, value=reg)
    # 订单数 = COUNTIFS（不含取消）
    ws2.cell(row=i, column=2).value = f'=COUNTIFS(订单明细!B2:B{last},A{i},订单明细!G2:G{last},"<>取消")'
    # 毛额 = SUMIFS
    ws2.cell(row=i, column=3).value = f'=SUMIFS(订单明细!H2:H{last},订单明细!B2:B{last},A{i})'
    ws2.cell(row=i, column=3).number_format = money_fmt
    # 退款 = SUMIFS
    ws2.cell(row=i, column=4).value = f'=SUMIFS(订单明细!I2:I{last},订单明细!B2:B{last},A{i})'
    ws2.cell(row=i, column=4).number_format = money_fmt
    # 净额 = 毛额 - 退款
    ws2.cell(row=i, column=5).value = f"=C{i}-D{i}"
    ws2.cell(row=i, column=5).number_format = money_fmt

# 总计行
tr2 = len(regions_order) + 2
ws2.cell(row=tr2, column=1, value="总计")
ws2.cell(row=tr2, column=1).font = Font(bold=True)
ws2.cell(row=tr2, column=2).value = f"=SUM(B2:B{tr2-1})"
ws2.cell(row=tr2, column=2).font = Font(bold=True)
for c in (3, 4, 5):
    ws2.cell(row=tr2, column=c).value = f"=SUM({chr(64+c)}2:{chr(64+c)}{tr2-1})"
    ws2.cell(row=tr2, column=c).font = Font(bold=True)
    ws2.cell(row=tr2, column=c).number_format = money_fmt

style_data(ws2, len(regions_order), len(headers2))
for col, w in zip("ABCDE", [8, 8, 10, 10, 10]):
    ws2.column_dimensions[col].width = w

# ── 保存 ─────────────────────────────────────────────────
wb.save(OUT)
print(f"已生成: {OUT}")
