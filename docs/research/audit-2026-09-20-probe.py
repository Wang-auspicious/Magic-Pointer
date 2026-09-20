"""Read-only product audit: synthetic files only, no GUI/provider calls.

Print observations rather than modifying implementation or asserting a fix.
Run from the repository with PYTHONPATH=. .
"""
from __future__ import annotations

import json
import tempfile
from dataclasses import asdict
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from pptx import Presentation

from app.actions.document_output import DocumentOutputHandler
from app.actions.figma import FigmaActionHandler
from app.actions.office_document import OfficeDocumentActionHandler
from app.actions.pdf import PdfActionHandler
from app.artifacts.document_patch import PatchOperation
from app.context_pack.sources import FragmentLocator, SourceRef


def source(path: Path, **identity) -> SourceRef:
    return SourceRef(
        source_id="audit-source", task_id="audit-task", kind="document",
        title=path.name, identity={"absolutePath": str(path), **identity},
        revision={}, capabilities=("read", "patch"), origin="user-attached",
        parent_source_id=None,
    )


def operation(name, kind, locator, before, after, op_id="audit-op"):
    return PatchOperation.from_dict({
        "operationId": op_id, "operation": name, "referenceId": "A",
        "sourceId": "audit-source", "locator": {"kind": kind, "value": locator},
        "before": before, "after": after,
    })


def main():
    observations = {}
    from app.ai_client import is_ai_failure
    from app.agent_runtime.compaction_prompt import summarize_history_text
    with patch("app.ai_client.get_ai_config", return_value=(None, "https://example.invalid/v1", "fake")), \
         patch("app.ai_client.get_ai_api_mode", return_value="chat_completions"), \
         patch("app.ai_client.record_unconfigured"):
        summary = summarize_history_text("Original task history: preserve the final correction.")
    observations["unconfigured_model_accepted_as_summary"] = {
        "summary": summary, "is_failure": is_ai_failure(summary),
        "nonempty_summary": bool(summary),
    }
    with tempfile.TemporaryDirectory(prefix="mp-audit-20260920-") as directory:
        root = Path(directory)
        word = root / "style.docx"
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("Keep ").bold = True
        paragraph.add_run("old").italic = True
        paragraph.add_run(" ending").underline = True
        doc.save(word)
        op = operation("replace_text", "text", {"paragraphIndex": 0},
                       {"text": "Keep old ending"}, {"text": "Keep new ending"})
        result = OfficeDocumentActionHandler().execute(source(word), op)
        paragraph = Document(word).paragraphs[0]
        observations["word_run_boundary"] = {
            "result": asdict(result),
            "runs": [{"text": r.text, "bold": r.bold, "italic": r.italic,
                      "underline": r.underline} for r in paragraph.runs],
        }

        original = root / "original.pdf"
        output = root / "annotated.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((40, 60), "first target; second target")
        doc.save(original)
        doc.close()
        handler = PdfActionHandler()
        def annotation(identity, rect):
            before = {"annotationId": identity, "kind": "text-note",
                      "outputPath": str(output), "present": False, "text": "wanted"}
            return operation("add_pdf_annotation", "pdf-region",
                             {"pageIndex": 0, "rectPt": rect}, before,
                             {**before, "present": True}, identity)
        first = annotation("first", [40, 40, 90, 60])
        second = annotation("second", [120, 40, 170, 60])
        first_result = handler.execute(source(original), first)
        second_pre_read = handler.read_current(source(original), second)
        second_result = handler.execute(source(original), second)
        observations["pdf_two_annotations_one_copy"] = {
            "first": asdict(first_result), "second_pre_read": asdict(second_pre_read),
            "second": asdict(second_result),
        }
        with fitz.open(output) as altered:
            page = altered[0]
            annot = next(page.annots())
            annot.set_info(content="different actual content")
            annot.update()
            altered.saveIncr()
        observations["pdf_readback_ignores_content"] = {
            "readback_equals_requested_after": handler.read_current(source(original), first).value == first.after,
            "actual_content": "different actual content",
        }

        class FakeFigma:
            task_id = "audit-task"
            document_session_id = "document"
            text = "abcDEFghi"
            def request(self, command, params):
                if command == "apply_patch":
                    change = params["operations"][0]
                    self.text = self.text[:change["start"]] + change["after"] + self.text[change["end"]:]
                    return {"ok": True}
                return {"nodes": [{"id": "node", "characters": self.text}]}
        client = FakeFigma()
        op = operation("replace_text", "figma-node",
                       {"nodeId": "node", "textStart": 3, "textEnd": 6}, "DEF", "LONGER")
        result = FigmaActionHandler(client).execute(source(root, documentSessionId="document"), op)
        observations["figma_length_changing_readback"] = {"actual": client.text, "result": asdict(result)}

        output_handler = DocumentOutputHandler()
        xlsx = root / "unnamed-sheet.xlsx"
        op = operation("create_file", "text", {"path": str(xlsx)},
                       {"exists": False, "path": str(xlsx)},
                       {"exists": True, "path": str(xlsx), "format": "xlsx",
                        "content": {"sheets": [{"rows": [["value"]]}]}, "references": []})
        observations["xlsx_supported_default_sheet_name"] = asdict(output_handler.execute(source(root), op))

        pptx = root / "tables.pptx"
        content = {"slides": [{"title": "Three tables", "body": "",
                               "tables": [{"rows": [["a", "b"], [1, 2]]} for _ in range(3)]}]}
        op = operation("create_file", "text", {"path": str(pptx)},
                       {"exists": False, "path": str(pptx)},
                       {"exists": True, "path": str(pptx), "format": "pptx",
                        "content": content, "references": []})
        result = output_handler.execute(source(root), op)
        deck = Presentation(pptx)
        tables = [shape for shape in deck.slides[0].shapes if shape.has_table]
        observations["pptx_verified_off_slide_table"] = {
            "result": asdict(result), "slide_height": deck.slide_height,
            "table_bottoms": [shape.top + shape.height for shape in tables],
            "off_slide": any(shape.top + shape.height > deck.slide_height for shape in tables),
        }
    print(json.dumps(observations, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
