"""Behavioral contracts for W08 document actions.

The fakes model the parts of Office identity that matter: two open windows can
show same-named presentations, slide order can change, and grouped shapes can
reuse display names.  Production code must bind by hwnd + full path and then
by slideId + shapeId, never by title, slide index, or shape name.
"""

from __future__ import annotations

from pathlib import Path
import os
import pytest

import fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.actions.office import make_word_replace_selection_proposal
from app.actions.office_document import OfficeDocumentActionHandler
from app.actions.document_backend import DocumentOperationBackend
from app.actions.document_output import DocumentOutputHandler
from app.actions.file_organizer import FileOrganizerHandler, restore_move
from app.actions.pdf import PdfActionHandler
from app.actions.powerpoint import PowerPointActionHandler
from app.artifacts.document_patch import PatchOperation
from app.context_pack.sources import FragmentLocator, SourceRef
from app.adapters.base import AdapterReadContext


@pytest.mark.skipif(os.name != 'nt', reason='Production binding executes Windows PowerShell')
def test_live_excel_binding_uses_the_selected_workbook_when_another_is_active() -> None:
    from app.actions.office_document import PowerShellLiveOfficeGateway
    from app.adapters.office_adapter import _run_powershell_json

    fixture = r'''
$sheet = [pscustomobject]@{}
$sheet | Add-Member ScriptMethod Range { param($address) return @{Address=$address} }
$sheets = [pscustomobject]@{}
$sheets | Add-Member ScriptMethod Item { param($name) return $sheet }
$target = [pscustomobject]@{FullName='D:\work\target.xlsx'; Worksheets=$sheets; Windows=@()}
$other = [pscustomobject]@{FullName='D:\work\other.xlsx'; Worksheets=$sheets; Windows=@()}
$application = [pscustomobject]@{ActiveWorkbook=$other; Workbooks=@($target,$other); Windows=@()}
$selected = [pscustomobject]@{HWND=101; Parent=$application}
$background = [pscustomobject]@{HWND=202; Parent=$application}
$target.Windows=@($selected); $other.Windows=@($background)
$application.Windows=@($selected,$background)
$p = [pscustomobject]@{hwnd=101; path='D:\work\target.xlsx'; sheet='Data'; address='A1'}
try {
'''
    script = fixture + PowerShellLiveOfficeGateway._EXCEL_BIND + r'''
  @{ok=$true; path=$workbook.FullName; hwnd=$window.HWND} | ConvertTo-Json -Compress
} catch { @{ok=$false; error=$_.Exception.Message} | ConvertTo-Json -Compress }
'''
    result = _run_powershell_json(script, timeout=12)
    assert result.ok, result.error
    assert result.data.get('ok') is True, result.data
    assert result.data['path'] == r'D:\work\target.xlsx'
    assert result.data['hwnd'] == 101


def _source(
    path: Path,
    *,
    source_id: str = "source-document",
    hwnd: int | None = None,
    kind: str = "document",
) -> SourceRef:
    identity = {"absolutePath": str(path.resolve())}
    if hwnd is not None:
        identity["hwnd"] = hwnd
    return SourceRef(
        source_id=source_id,
        task_id="task-office",
        kind=kind,
        title=path.name,
        identity=identity,
        revision={"authority": "test"},
        capabilities=("read", "patch"),
        origin="user-attached",
        parent_source_id=None,
    )


def _operation(
    name: str,
    locator: FragmentLocator,
    before,
    after,
    *,
    source_id: str = "source-document",
) -> PatchOperation:
    return PatchOperation.from_dict({
        "operationId": f"operation-{name}",
        "operation": name,
        "referenceId": "reference-target",
        "sourceId": source_id,
        "locator": locator.to_dict(),
        "before": before,
        "after": after,
    })


class FakeShape:
    def __init__(self, shape_id: int, name: str, text: str, children=()) -> None:
        self.shape_id = shape_id
        self.name = name
        self.text = text
        self.children = list(children)
        self.geometry = {"left": 10.0, "top": 20.0, "width": 200.0, "height": 60.0}
        self.style = {"fillColor": "#FFFFFF", "lineColor": "#000000"}


class FakeSlide:
    def __init__(self, slide_id: int, shapes: list[FakeShape]) -> None:
        self.slide_id = slide_id
        self.shapes = shapes


class FakePresentation:
    def __init__(self, path: Path, hwnd: int, slides: list[FakeSlide]) -> None:
        self.path = str(path.resolve())
        self.hwnd = hwnd
        self.slides = slides
        self.saved = True


class FakePowerPointGateway:
    used_backend = "fake-powerpoint-com"

    def __init__(self, presentations: list[FakePresentation]) -> None:
        self.presentations = presentations
        self.writes: list[tuple[str, int, int]] = []

    def _presentation(self, path: str, hwnd: int) -> FakePresentation:
        return next(
            item for item in self.presentations
            if item.path.casefold() == str(Path(path).resolve()).casefold() and item.hwnd == hwnd
        )

    @staticmethod
    def _shape(shapes: list[FakeShape], shape_id: int) -> FakeShape:
        for shape in shapes:
            if shape.shape_id == shape_id:
                return shape
            try:
                return FakePowerPointGateway._shape(shape.children, shape_id)
            except LookupError:
                pass
        raise LookupError(shape_id)

    def read_shape(self, *, path: str, hwnd: int, slide_id: int, shape_id: int):
        presentation = self._presentation(path, hwnd)
        slide = next(item for item in presentation.slides if item.slide_id == slide_id)
        shape = self._shape(slide.shapes, shape_id)
        return {
            "text": shape.text,
            "geometry": dict(shape.geometry),
            "style": dict(shape.style),
            "locked": False,
            "masterShape": False,
            "presentationSaved": presentation.saved,
        }

    def replace_shape_text(
        self,
        *,
        path: str,
        hwnd: int,
        slide_id: int,
        shape_id: int,
        expected_text: str,
        start: int,
        length: int,
        replacement: str,
    ):
        presentation = self._presentation(path, hwnd)
        slide = next(item for item in presentation.slides if item.slide_id == slide_id)
        shape = self._shape(slide.shapes, shape_id)
        if shape.text != expected_text:
            return {"ok": False, "wrote": False, "error": "base_mismatch"}
        shape.text = shape.text[:start] + replacement + shape.text[start + length:]
        presentation.saved = False
        self.writes.append((path, slide_id, shape_id))
        return {"ok": True, "wrote": True}


def test_powerpoint_uses_full_path_hwnd_slide_id_and_recursive_shape_id(tmp_path: Path) -> None:
    right = tmp_path / "right" / "same-name.pptx"
    wrong = tmp_path / "wrong" / "same-name.pptx"
    right.parent.mkdir()
    wrong.parent.mkdir()
    target = FakeShape(42, "Body", "Keep this concise")
    gateway = FakePowerPointGateway([
        FakePresentation(wrong, 101, [FakeSlide(700, [FakeShape(42, "Body", "wrong")])]),
        FakePresentation(
            right,
            202,
            # Slide 700 is deliberately second: slide order is not identity.
            [FakeSlide(999, [FakeShape(42, "Body", "other slide")]),
             FakeSlide(700, [FakeShape(9, "Group", "", [target])])],
        ),
    ])
    source = _source(right, hwnd=202)
    locator = FragmentLocator("slide-shape", {"slideId": 700, "shapeId": 42})
    operation = _operation(
        "set_shape_text",
        locator,
        {"text": "Keep this concise"},
        {"text": "Keep concise"},
    )
    handler = PowerPointActionHandler(gateway=gateway)

    assert handler.read_current(source, operation).value == operation.before
    result = handler.execute(source, operation)

    assert result.ok is True and result.wrote is True
    assert target.text == "Keep concise"
    assert gateway.writes == [(str(right.resolve()), 700, 42)]


def test_powerpoint_changed_base_never_writes(tmp_path: Path) -> None:
    path = tmp_path / "deck.pptx"
    target = FakeShape(42, "Body", "user changed this")
    gateway = FakePowerPointGateway([
        FakePresentation(path, 202, [FakeSlide(700, [target])]),
    ])
    operation = _operation(
        "set_shape_text",
        FragmentLocator("slide-shape", {"slideId": 700, "shapeId": 42}),
        {"text": "old"},
        {"text": "new"},
    )
    handler = PowerPointActionHandler(gateway=gateway)

    assert handler.read_current(_source(path, hwnd=202), operation).value == {
        "text": "user changed this",
    }
    assert gateway.writes == []


def test_powerpoint_rejects_unknown_style_and_invalid_geometry_fields(tmp_path: Path) -> None:
    path = tmp_path / "deck.pptx"
    target = FakeShape(42, "Body", "text")
    target.style = {"fillRgb": 0xFFFFFF, "lineRgb": 0}
    gateway = FakePowerPointGateway([
        FakePresentation(path, 202, [FakeSlide(700, [target])]),
    ])
    handler = PowerPointActionHandler(gateway=gateway)
    locator = FragmentLocator("slide-shape", {"slideId": 700, "shapeId": 42})

    unknown_style = _operation(
        "set_shape_style",
        locator,
        {"shadowBlur": 0},
        {"shadowBlur": 20},
    )
    style_read = handler.read_current(_source(path, hwnd=202), unknown_style)
    assert style_read.ok is False
    assert "unsupported" in str(style_read.error)

    invalid_geometry = _operation(
        "set_shape_geometry",
        locator,
        {"width": 200.0},
        {"width": -5.0},
    )
    geometry_write = handler.execute(_source(path, hwnd=202), invalid_geometry)
    assert geometry_write.ok is False and geometry_write.wrote is False
    assert "positive" in str(geometry_write.error)
    assert gateway.writes == []


def test_pdf_rotated_highlight_is_written_to_new_copy_only(tmp_path: Path) -> None:
    original = tmp_path / "source.pdf"
    output = tmp_path / "annotated.pdf"
    document = fitz.open()
    page = document.new_page(width=300, height=200)
    page.insert_text((40, 60), "rotation target")
    page.set_rotation(90)
    document.save(original)
    document.close()
    original_bytes = original.read_bytes()
    locator = FragmentLocator("pdf-region", {
        "pageIndex": 0,
        "rectPt": [20.0, 30.0, 90.0, 55.0],
        "coordinateSpace": "rotated-page-points",
    })
    before = {
        "annotationId": "annotation-rotated",
        "kind": "highlight",
        "outputPath": str(output),
        "present": False,
    }
    after = {**before, "present": True}
    operation = _operation("add_pdf_annotation", locator, before, after)
    handler = PdfActionHandler()

    assert handler.read_current(_source(original), operation).value == before
    result = handler.execute(_source(original), operation)

    assert result.ok is True and result.wrote is True
    assert original.read_bytes() == original_bytes
    assert output.exists()
    reopened = fitz.open(output)
    reopened_page = reopened[0]
    annotations = list(reopened_page.annots() or ())
    assert len(annotations) == 1
    assert annotations[0].info["subject"] == "Magic Pointer annotation-rotated"
    # The visual-space rectangle is converted through the page's de-rotation
    # matrix, so a 90-degree page must not receive the raw rectangle.
    assert tuple(round(value, 2) for value in annotations[0].rect) != (20.0, 30.0, 90.0, 55.0)
    reopened.close()


def test_create_docx_xlsx_and_pptx_are_reopened_and_verified(tmp_path: Path) -> None:
    handler = DocumentOutputHandler()
    cases = [
        (
            "docx",
            tmp_path / "report.docx",
            {"paragraphs": ["Executive summary", "Evidence-backed conclusion"],
             "tables": [{"rows": [["Metric", "Value"], ["Coverage", "100%"]]}]},
        ),
        (
            "xlsx",
            tmp_path / "report.xlsx",
            {"sheets": [{"name": "Summary", "rows": [["Metric", "Value"], ["Total", "=SUM(1,2)"]]}]},
        ),
        (
            "pptx",
            tmp_path / "report.pptx",
            {"slides": [{"title": "Executive summary", "body": "Evidence-backed conclusion"}]},
        ),
    ]
    for file_format, path, content in cases:
        before = {"exists": False, "path": str(path)}
        after = {
            "exists": True,
            "path": str(path),
            "format": file_format,
            "content": content,
            "references": [{"sourceId": "source-material", "label": "Input brief"}],
        }
        operation = _operation(
            "create_file",
            FragmentLocator("text", {"path": str(path)}),
            before,
            after,
        )
        result = handler.execute(_source(tmp_path, kind="file"), operation)
        assert result.ok is True and result.wrote is True
        assert handler.read_current(_source(tmp_path, kind="file"), operation).value == after

    assert Document(tmp_path / "report.docx").paragraphs[0].text == "Executive summary"
    assert load_workbook(tmp_path / "report.xlsx", data_only=False)["Summary"]["B2"].value == "=SUM(1,2)"
    assert len(Presentation(tmp_path / "report.pptx").slides) >= 2  # content + editable sources slide


def test_file_move_conflict_does_not_overwrite_and_inverse_restores(tmp_path: Path) -> None:
    old_path = tmp_path / "inbox" / "brief.txt"
    new_path = tmp_path / "archive" / "brief.txt"
    old_path.parent.mkdir()
    new_path.parent.mkdir()
    old_path.write_text("source", encoding="utf-8")
    new_path.write_text("existing", encoding="utf-8")
    before = {"path": str(old_path)}
    after = {"path": str(new_path)}
    operation = _operation(
        "move_file",
        FragmentLocator("text", {"root": str(tmp_path)}),
        before,
        after,
    )
    handler = FileOrganizerHandler()

    conflict = handler.execute(_source(tmp_path, kind="file"), operation)
    assert conflict.ok is False and conflict.wrote is False
    assert old_path.read_text(encoding="utf-8") == "source"
    assert new_path.read_text(encoding="utf-8") == "existing"

    new_path.unlink()
    moved = handler.execute(_source(tmp_path, kind="file"), operation)
    assert moved.ok is True and new_path.read_text(encoding="utf-8") == "source"
    restored = restore_move(expected_current=new_path, restore=old_path)
    assert restored.ok is True
    assert old_path.read_text(encoding="utf-8") == "source"
    assert not new_path.exists()


def test_document_backend_executes_through_safe_action_executor(tmp_path: Path) -> None:
    path = tmp_path / "output.docx"
    source = _source(tmp_path, kind="file")
    operation = _operation(
        "create_file",
        FragmentLocator("text", {"path": str(path)}),
        {"exists": False, "path": str(path)},
        {
            "exists": True,
            "path": str(path),
            "format": "docx",
            "content": {"paragraphs": ["Created through the shared executor"]},
            "references": [],
        },
    )
    backend = DocumentOperationBackend(
        sources=(source,),
        artifact_id="artifact-1",
        artifact_revision=2,
    )

    result = backend.execute(operation)

    assert result.ok is True and result.wrote is True
    assert "safe-action-executor" in result.used_backend
    assert path.exists()


def test_word_proposal_carries_source_locator_base_and_artifact_binding() -> None:
    locator = FragmentLocator("text", {"story": "selection", "start": 12, "end": 18})
    context = AdapterReadContext(
        adapter="office",
        app="word",
        label=r"D:\Docs\resume.docx",
        method="com:word.selection",
        content="before",
        window={"title": "resume.docx - Word"},
        artifacts={
            "document": r"D:\Docs\resume.docx",
            "document_name": "resume.docx",
            "hwnd": 88,
            "selection_start": 12,
            "selection_end": 18,
        },
    )

    proposal = make_word_replace_selection_proposal(
        context,
        command="Shorten only this selection",
        replacement_text="after",
        source_id="source-word",
        locator=locator,
        artifact_id="artifact-word",
        artifact_revision=3,
    )

    assert proposal is not None
    assert proposal.parameters["source_id"] == "source-word"
    assert proposal.parameters["locator"] == locator.to_dict()
    assert proposal.parameters["base"] == {"text": "before"}
    assert proposal.parameters["artifact_id"] == "artifact-word"
    assert proposal.parameters["artifact_revision"] == 3


def test_docx_range_replacement_preserves_unselected_run_formatting(tmp_path: Path) -> None:
    path = tmp_path / "resume.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Keep ").bold = True
    paragraph.add_run("old").italic = True
    paragraph.add_run(" ending").underline = True
    document.save(path)
    operation = _operation(
        "replace_text",
        FragmentLocator("text", {"paragraphIndex": 0}),
        {"text": "Keep old ending"},
        {"text": "Keep new ending"},
    )
    handler = OfficeDocumentActionHandler()

    result = handler.execute(_source(path), operation)

    assert result.ok is True
    reopened = Document(path)
    assert reopened.paragraphs[0].text == "Keep new ending"
    assert reopened.paragraphs[0].runs[0].bold is True
    assert reopened.paragraphs[0].runs[-1].underline is True
    assert handler.read_current(_source(path), operation).value == operation.after


def test_xlsx_cell_patch_reads_and_writes_formulas_without_claiming_cache(tmp_path: Path) -> None:
    path = tmp_path / "metrics.xlsx"
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "untouched"
    sheet["B2"] = 3
    sheet["B3"] = "=SUM(1,2)"
    workbook.save(path)
    operation = _operation(
        "set_cell_values",
        FragmentLocator("cell-range", {"sheet": "Summary", "range": "B2:B3"}),
        [[3], ["=SUM(1,2)"]],
        [[4], ["=SUM(2,3)"]],
    )
    handler = OfficeDocumentActionHandler()

    assert handler.read_current(_source(path), operation).value == operation.before
    result = handler.execute(_source(path), operation)

    assert result.ok is True
    reopened = load_workbook(path, data_only=False)
    assert reopened["Summary"]["A1"].value == "untouched"
    assert reopened["Summary"]["B2"].value == 4
    assert reopened["Summary"]["B3"].value == "=SUM(2,3)"
    reopened.close()
    assert handler.read_current(_source(path), operation).value == operation.after


class FakeLiveOfficeGateway:
    used_backend = "fake-live-office-com"

    def __init__(self) -> None:
        self.word_text = "selected old text"
        self.calls: list[tuple] = []

    def read_word(self, *, path: str, hwnd: int, start: int, end: int):
        self.calls.append(("read-word", path, hwnd, start, end))
        return self.word_text

    def replace_word(
        self,
        *,
        path: str,
        hwnd: int,
        start: int,
        end: int,
        expected_text: str,
        replacement: str,
    ):
        self.calls.append(("write-word", path, hwnd, start, end))
        if self.word_text != expected_text:
            return {"ok": False, "wrote": False, "error": "base_mismatch"}
        self.word_text = replacement
        return {"ok": True, "wrote": True}


def test_live_word_patch_is_bound_to_full_path_hwnd_and_range(tmp_path: Path) -> None:
    path = tmp_path / "same-name.docx"
    gateway = FakeLiveOfficeGateway()
    handler = OfficeDocumentActionHandler(gateway=gateway)
    source = _source(path, hwnd=808)
    operation = _operation(
        "replace_text",
        FragmentLocator("text", {"story": "selection", "start": 14, "end": 31}),
        {"text": "selected old text"},
        {"text": "selected new text"},
    )

    assert handler.read_current(source, operation).value == operation.before
    result = handler.execute(source, operation)

    assert result.ok is True
    assert gateway.word_text == "selected new text"
    assert ("write-word", str(path.resolve()), 808, 14, 31) in gateway.calls
