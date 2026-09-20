"""Regressions for source continuity found by tracing the production readers."""

from pathlib import Path

import pytest
from docx import Document

from app.context_pack.document_reader import DocumentReader
from app.context_pack.selection_reader import FrozenSelectionMaterial, FrozenSelectionReader
from app.context_pack.sources import Coverage, FragmentLocator, SourceRef


def source_for(path: Path) -> SourceRef:
    return SourceRef(
        source_id="source:audit", task_id="task-audit", kind="file", title=path.name,
        identity={"absolutePath": str(path)}, revision={},
        capabilities=("read", "search", "follow"), origin="user-attached", parent_source_id=None,
    )


@pytest.mark.parametrize("limit", [1, 2])
def test_frozen_search_pagination_does_not_skip_disk_hits(tmp_path, limit):
    path = tmp_path / "notes.txt"
    path.write_text("match first\nmatch second\nmatch third", encoding="utf-8")
    source = source_for(path)
    material = FrozenSelectionMaterial(
        source.source_id, "match unsaved selection", FragmentLocator("text", {"selection": True}),
        Coverage("selection", (), 1, True, None, None), "frozen.selection",
    )
    reader = FrozenSelectionReader((material,), fallback=DocumentReader())
    cursor = None
    texts = []
    for _ in range(8):
        result = reader.search(source, "match", cursor, limit)
        assert len(result.fragments) <= limit
        texts.extend(fragment.text for fragment in result.fragments)
        cursor = result.coverage.next_cursor
        if cursor is None:
            assert result.coverage.complete
            break
    assert texts == ["match unsaved selection", "match first", "match second", "match third"]


def test_frozen_only_reader_does_not_relabel_selection_as_an_unseen_page(tmp_path):
    source = source_for(tmp_path / "capture.txt")
    original = FragmentLocator("pdf-region", {"pageIndex": 0})
    material = FrozenSelectionMaterial(
        source.source_id, "page one", original,
        Coverage("selection", ({"pageIndex": 0},), 1, True, None, None), "frozen.selection",
    )
    result = FrozenSelectionReader((material,)).read(
        source, FragmentLocator("pdf-region", {"pageIndex": 36}), None, 8,
    )
    assert result.fragments == ()
    assert result.evidence_status == "unsupported"
    assert not result.coverage.complete


def test_follow_directory_hit_survives_an_earlier_filename_being_added(tmp_path):
    folder = tmp_path / "materials"
    folder.mkdir()
    target = folder / "b-contract.txt"
    target.write_text("chosen liability clause", encoding="utf-8")
    source = source_for(folder)
    reader = DocumentReader()
    result = reader.search(source, "liability", None, 8)
    (folder / "a-unrelated.txt").write_text("unrelated", encoding="utf-8")
    followed = reader.follow(source, result.fragments[0].fragment_id)
    assert len(followed) == 1
    assert Path(followed[0].identity["absolutePath"]) == target


def test_removed_directory_hit_does_not_follow_its_replacement_at_same_index(tmp_path):
    folder = tmp_path / "materials"
    folder.mkdir()
    target = folder / "a-contract.txt"
    target.write_text("chosen liability clause", encoding="utf-8")
    source = source_for(folder)
    reader = DocumentReader()
    result = reader.search(source, "liability", None, 8)
    target.unlink()
    (folder / "b-unrelated.txt").write_text("unrelated", encoding="utf-8")
    assert reader.follow(source, result.fragments[0].fragment_id) == ()


def test_precise_document_locator_returns_target_before_neighborhood(tmp_path):
    path = tmp_path / "conditions.docx"
    document = Document()
    document.add_paragraph("Earlier condition to preserve")
    document.add_paragraph("Selected clause to edit")
    document.add_paragraph("Later condition to preserve")
    document.save(path)
    source = source_for(path)
    reader = DocumentReader()
    target = reader.search(source, "Selected clause", None, 1).fragments[0]
    result = reader.read(source, target.locator, None, 1)
    assert result.fragments[0].text == target.text
    assert result.fragments[0].locator == target.locator
    remaining = reader.read(source, target.locator, result.coverage.next_cursor, 8)
    assert [item.text for item in remaining.fragments] == [
        "Earlier condition to preserve", "Later condition to preserve",
    ]


def test_context_bind_labels_continue_after_z(tmp_path):
    from app.agent_runtime.session import FileSessionStore
    from app.agent_runtime.tool_registry import ToolRegistry
    from app.context_pack.source_store import register_source, apply_reference_updates, task_references
    from app.context_pack.sources import ReferenceBinding, ReferenceUpdate, SourceReaderRegistry
    from app.context_pack.tools import register_context_tools

    path = tmp_path / "references.txt"
    path.write_text("selected sentence", encoding="utf-8")
    source = source_for(path)
    session = FileSessionStore(tmp_path / "sessions").open_or_create(source.task_id)
    register_source(session, source)
    locator = FragmentLocator("text", {"lineStart": 1, "lineEnd": 1})
    updates = tuple(ReferenceUpdate("add", ReferenceBinding(
        f"reference:{i}", chr(64 + i), source.source_id, locator,
        "reference", None, 1, i, True,
    )) for i in range(1, 27))
    apply_reference_updates(session, updates)
    readers = SourceReaderRegistry()
    readers.register("file", DocumentReader())
    registry = ToolRegistry()
    register_context_tools(registry, session=session, readers=readers)
    assert not registry.execute_tool("Context.read", {"source_id": source.source_id}).is_error
    bound = registry.execute_tool("Context.bind", {
        "reference_id": "reference:new", "source_id": source.source_id,
        "locator": locator.to_dict(), "role": "target", "reason": "current clause",
    })
    assert not bound.is_error, bound
    assert task_references(session.events)[-1].label == "AA"


def test_long_document_unit_is_fully_reachable_through_read_cursors(tmp_path):
    path = tmp_path / "long-record.json"
    content = '{"record":"' + "long content " * 3_000 + '"}'
    path.write_text(content, encoding="utf-8")
    source = source_for(path)
    reader = DocumentReader()
    first = reader.read(source, None, None, 1)
    assert not first.coverage.complete
    assert first.coverage.next_cursor
    pieces = list(first.fragments)
    cursor = first.coverage.next_cursor
    while cursor:
        result = reader.read(source, None, cursor, 1)
        pieces.extend(result.fragments)
        cursor = result.coverage.next_cursor
    assert "".join(fragment.text for fragment in pieces) == content
    reread = reader.read(source, pieces[1].locator, None, 1)
    assert reread.fragments[0].text == pieces[1].text
