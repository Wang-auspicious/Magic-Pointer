from __future__ import annotations

import json
from pathlib import Path

import fitz
import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from app.adapters import office_adapter
from app.context_pack.document_reader import DocumentReader
from app.context_pack.selection_reader import FrozenSelectionMaterial, FrozenSelectionReader
from app.context_pack.sources import Coverage, FragmentLocator, SourceReaderRegistry, SourceRef
from app.file_context import read_local_file_context


def _source(path: Path, *, source_id: str = "source:file", kind: str = "document") -> SourceRef:
    return SourceRef(
        source_id=source_id,
        task_id="task-document-reader",
        kind=kind,
        title=path.name,
        identity={"absolutePath": str(path.resolve())},
        revision={"mtimeNs": path.stat().st_mtime_ns},
        capabilities=("read", "search", "follow"),
        origin="user-attached",
        parent_source_id=None,
    )


@pytest.fixture()
def pdf_path(tmp_path: Path) -> Path:
    path = tmp_path / "forty-pages.pdf"
    document = fitz.open()
    for index in range(40):
        page = document.new_page(width=595, height=842)
        text = f"Ordinary material on page {index + 1}."
        if index == 36:
            text += " UNIQUE LIABILITY LIMIT: supplier liability is capped at 37 units."
        page.insert_text((72, 96), text, fontsize=11)
    document.save(path)
    document.close()
    return path


@pytest.fixture()
def docx_path(tmp_path: Path) -> Path:
    path = tmp_path / "ordered-body.docx"
    document = Document()
    document.add_heading("Decision record", level=1)
    document.add_paragraph("Condition before the table.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Decision"
    table.cell(0, 1).text = "Owner"
    table.cell(1, 0).text = "FINAL TABLE CONCLUSION"
    table.cell(1, 1).text = "Ada"
    document.add_paragraph("Constraint after the table.")
    section = document.sections[0]
    section.header.paragraphs[0].text = "Confidential header"
    section.footer.paragraphs[0].text = "Approval footer"
    document.save(path)
    return path


@pytest.fixture()
def pptx_path(tmp_path: Path) -> Path:
    path = tmp_path / "shape-identity.pptx"
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    first = presentation.slides.add_slide(blank)
    first_box = first.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    first_box.name = "Repeated Name"
    first_box.text = "First slide duplicate"
    second = presentation.slides.add_slide(blank)
    second_box = second.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    second_box.name = "Repeated Name"
    second_box.text = "Second slide duplicate"
    group = second.shapes.add_group_shape()
    group.name = "Evidence Group"
    nested = group.shapes.add_textbox(Inches(2), Inches(3), Inches(4), Inches(1))
    nested.name = "Nested Target"
    nested.text = "GROUP-ONLY TARGET EVIDENCE"
    presentation.save(path)
    return path


@pytest.fixture()
def xlsx_path(tmp_path: Path) -> Path:
    path = tmp_path / "merged-header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Forecast"
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Revenue（万元）"
    sheet["A2"] = "Product"
    sheet["B2"] = "2026"
    sheet["A3"] = "Atlas"
    sheet["B3"] = 125
    sheet["A4"] = "Projected"
    sheet["B4"] = "=B3*2"
    sheet.row_dimensions[4].hidden = True
    hidden = workbook.create_sheet("Hidden assumptions")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "private multiplier"
    workbook.save(path)
    return path


def test_pdf_search_reaches_unique_clause_on_page_37(pdf_path: Path) -> None:
    reader = DocumentReader()
    source = _source(pdf_path)

    result = reader.search(source, "UNIQUE LIABILITY LIMIT", None, 5)

    assert result.evidence_status == "ok"
    assert result.used_backend == "document.pdf.pymupdf"
    assert result.coverage.extent == "query-results"
    assert result.coverage.total_units == 40
    assert result.coverage.complete is True
    assert len(result.fragments) == 1
    locator = result.fragments[0].locator
    assert locator.kind == "pdf-region"
    assert locator.value["pageIndex"] == 36
    assert len(locator.value["rectPt"]) == 4

    neighborhood = reader.read(source, locator, None, 5)
    assert "capped at 37 units" in "\n".join(item.text for item in neighborhood.fragments)
    assert neighborhood.coverage.extent == "neighborhood"


def test_pdf_search_falls_back_to_ranked_terms_when_mixed_language_phrase_is_absent(
    pdf_path: Path,
) -> None:
    reader = DocumentReader()
    source = _source(pdf_path)

    result = reader.search(
        source,
        "责任限制 liability limitation cap",
        None,
        5,
    )

    assert result.evidence_status == "ok"
    assert len(result.fragments) == 1
    assert result.fragments[0].locator.value["pageIndex"] == 36
    assert "capped at 37 units" in result.fragments[0].text


def test_dense_pdf_read_pages_before_runtime_would_truncate_the_middle(
    tmp_path: Path,
) -> None:
    path = tmp_path / "dense.pdf"
    document = fitz.open()
    for page_index in range(30):
        page = document.new_page(width=595, height=842)
        for block_index in range(5):
            marker = f"PAGE-{page_index + 1:02}-BLOCK-{block_index + 1}"
            text = marker + " " + ("material evidence " * 24)
            top = 48 + block_index * 150
            page.insert_textbox(
                fitz.Rect(36, top, 559, top + 130),
                text,
                fontsize=8,
            )
    document.save(path)
    document.close()
    reader = DocumentReader()
    source = _source(path)

    result = reader.read(source, None, "unit:50", 100)
    serialized = json.dumps(result.to_dict(), ensure_ascii=False)

    assert len(serialized) < 60_000
    assert result.coverage.complete is False
    assert result.coverage.next_cursor is not None
    last_index = int(result.fragments[-1].fragment_id.rsplit(":", 1)[-1])
    assert result.coverage.next_cursor == f"unit:{last_index + 1}"
    continued = reader.read(source, None, result.coverage.next_cursor, 1)
    assert continued.fragments[0].fragment_id.endswith(f":unit:{last_index + 1}")


def test_pdf_term_fallback_ranks_a_whole_page_not_isolated_text_blocks(
    tmp_path: Path,
) -> None:
    path = tmp_path / "split-clause.pdf"
    document = fitz.open()
    for page_index in range(40):
        page = document.new_page(width=595, height=842)
        page.insert_text(
            (72, 96),
            f"Section {page_index + 1} - Supporting Terms",
            fontsize=11,
        )
        page.insert_text(
            (72, 144),
            "No alternate aggregate liability cap is stated on this page.",
            fontsize=11,
        )
        if page_index == 36:
            page.insert_text(
                (72, 216),
                "SECTION 12.2 - LIMITATION OF LIABILITY",
                fontsize=11,
            )
            page.insert_text(
                (72, 264),
                "The aggregate cap and maximum recovery are SGD 100,000.",
                fontsize=11,
            )
    document.save(path)
    document.close()
    reader = DocumentReader()

    result = reader.search(
        _source(path),
        "liability cap limitation aggregate maximum",
        None,
        1,
    )

    assert result.fragments
    assert len(result.fragments) == 1
    assert {fragment.locator.value["pageIndex"] for fragment in result.fragments} == {36}
    assert "SGD 100,000" in "\n".join(fragment.text for fragment in result.fragments)
    tied_terms = reader.search(
        _source(path),
        "liability cap aggregate limit limitation",
        None,
        1,
    )
    assert tied_terms.fragments[0].locator.value["pageIndex"] == 36
    coarse_page_locator = FragmentLocator(
        "pdf-region",
        {"pageIndex": 36, "rectPt": [36.0, 80.0, 559.0, 300.0]},
    )
    neighborhood = reader.read(
        source := _source(path),
        coarse_page_locator,
        None,
        100,
    )
    neighborhood_pages = [
        fragment.locator.value["pageIndex"] for fragment in neighborhood.fragments
    ]
    assert neighborhood_pages[0] == 36
    assert set(neighborhood_pages) == {35, 36, 37}


def test_docx_preserves_paragraph_table_order_and_headers(docx_path: Path) -> None:
    result = DocumentReader().read(_source(docx_path), None, None, 20)
    texts = [fragment.text for fragment in result.fragments]

    before = texts.index("Condition before the table.")
    table = next(index for index, text in enumerate(texts) if "FINAL TABLE CONCLUSION" in text)
    after = texts.index("Constraint after the table.")
    assert before < table < after
    table_fragment = result.fragments[table]
    assert table_fragment.locator.kind == "table"
    assert table_fragment.locator.value["tableIndex"] == 0
    assert table_fragment.locator.value["rowIndex"] == 1
    assert any(fragment.metadata.get("story") == "header" for fragment in result.fragments)
    assert any(fragment.metadata.get("story") == "footer" for fragment in result.fragments)
    assert result.coverage.complete is True
    assert "floating-objects-and-revisions-not-expanded" in (result.coverage.missing_reason or "")


def test_pptx_uses_slide_and_shape_identity_and_recurses_groups(pptx_path: Path) -> None:
    reader = DocumentReader()
    source = _source(pptx_path)

    duplicates = reader.search(source, "duplicate", None, 10)
    duplicate_locators = [fragment.locator.value for fragment in duplicates.fragments]
    assert len(duplicate_locators) == 2
    assert duplicate_locators[0]["shapeName"] == duplicate_locators[1]["shapeName"] == "Repeated Name"
    assert duplicate_locators[0]["slideId"] != duplicate_locators[1]["slideId"]
    assert duplicate_locators[0]["shapeId"] == duplicate_locators[1]["shapeId"]

    nested = reader.search(source, "GROUP-ONLY TARGET", None, 10)
    assert len(nested.fragments) == 1
    locator = nested.fragments[0].locator
    assert locator.kind == "slide-shape"
    assert locator.value["parentShapeId"] is not None
    assert locator.value["shapeName"] == "Nested Target"
    assert len(locator.value["bboxPt"]) == 4
    assert nested.coverage.total_units == 2


def test_xlsx_value_keeps_merged_header_unit_formula_and_hidden_state(xlsx_path: Path) -> None:
    reader = DocumentReader()
    source = _source(xlsx_path)

    value_result = reader.search(source, "125", None, 10)
    assert len(value_result.fragments) == 1
    value = value_result.fragments[0]
    assert value.locator.kind == "cell-range"
    assert value.locator.value == {"sheet": "Forecast", "range": "B3"}
    assert value.metadata["headers"] == ["Revenue（万元）", "2026"]
    assert value.metadata["unit"] == "万元"
    assert "unit=万元" in value.text

    formula_result = reader.search(source, "B3*2", None, 10)
    formula = formula_result.fragments[0]
    assert formula.metadata["formula"] == "=B3*2"
    assert formula.metadata["cachedValueKnown"] is False
    assert formula.metadata["rowHidden"] is True

    hidden_result = reader.search(source, "private multiplier", None, 10)
    assert hidden_result.fragments[0].metadata["sheetHidden"] is True


def test_directory_and_text_reads_page_instead_of_stopping_at_120(tmp_path: Path) -> None:
    folder = tmp_path / "many"
    folder.mkdir()
    for index in range(125):
        (folder / f"item-{index:03}.txt").write_text(f"value {index}", encoding="utf-8")
    source = _source(folder, kind="file")
    reader = DocumentReader()

    first = reader.read(source, None, None, 120)
    assert 0 < len(first.fragments) < 120
    assert first.coverage.complete is False
    assert first.coverage.next_cursor == f"unit:{len(first.fragments)}"
    fragments = list(first.fragments)
    cursor = first.coverage.next_cursor
    while cursor is not None:
        page = reader.read(source, None, cursor, 120)
        fragments.extend(page.fragments)
        cursor = page.coverage.next_cursor
    assert len(fragments) == 125
    assert fragments[-1].text.endswith("item-124.txt")

    content_match = reader.search(source, "value 124", None, 5)
    assert len(content_match.fragments) == 1
    assert content_match.fragments[0].metadata["relativePath"] == "item-124.txt"
    assert content_match.fragments[0].metadata["matchedBackend"].startswith("document.text")
    assert content_match.fragments[0].metadata["matchedChildLocator"]["value"]["lineStart"] == 1
    followed = reader.follow(source, content_match.fragments[0].fragment_id)
    assert len(followed) == 1
    assert followed[0].identity["absolutePath"] == str((folder / "item-124.txt").resolve())


def test_unknown_format_is_unsupported_not_empty(tmp_path: Path) -> None:
    path = tmp_path / "opaque.bin"
    path.write_bytes(b"\x00\x01\x02")

    result = DocumentReader().read(_source(path, kind="file"), None, None, 5)

    assert result.evidence_status == "unsupported"
    assert result.fragments == ()
    assert result.coverage.complete is False
    assert result.coverage.missing_reason == "unsupported-file-type:.bin"


def test_source_specific_live_reader_wins_over_stale_disk_reader(tmp_path: Path) -> None:
    path = tmp_path / "draft.txt"
    path.write_text("old disk content", encoding="utf-8")
    source = _source(path, source_id="source:live", kind="document")
    locator = FragmentLocator("text", {"story": "selection", "index": 0})
    live = FrozenSelectionReader((FrozenSelectionMaterial(
        source_id=source.source_id,
        text="new unsaved live content",
        locator=locator,
        coverage=Coverage("selection", ({"story": "selection"},), 1, True, None, None),
        used_backend="office.com.live",
    ),))
    readers = SourceReaderRegistry()
    readers.register("document", DocumentReader())
    readers.register_source(source.source_id, live)

    result = readers.for_source(source).read(source, None, None, 5)

    assert result.fragments[0].text == "new unsaved live content"
    assert result.used_backend == "office.com.live"


def test_file_context_preview_exposes_source_coverage_and_structure(pdf_path: Path) -> None:
    context = read_local_file_context(str(pdf_path), max_chars=1_000)

    assert context.source_id
    assert context.coverage["totalUnits"] == 40
    assert context.structure["pageCount"] == 40
    assert context.method == "document.pdf.pymupdf"
    assert context.truncated is True


def test_powerpoint_native_reader_binds_requested_hwnd_and_returns_shape_locator(monkeypatch) -> None:
    captured: dict[str, str] = {}

    def fake_probe(script: str, *, timeout: int = 2) -> office_adapter.OfficeProbeResult:
        captured["script"] = script
        return office_adapter.OfficeProbeResult(True, {
            "app": "powerpoint",
            "method": "com:powerpoint.selection",
            "hwnd": 4242,
            "presentation": "D:\\Decks\\live.pptx",
            "presentation_name": "live.pptx",
            "presentation_saved": False,
            "slide_id": 512,
            "slide_index": 2,
            "shapes": [{
                "shape_id": 7,
                "name": "Target",
                "type": 17,
                "left": 72,
                "top": 144,
                "width": 216,
                "height": 72,
                "text": "Unsaved slide text",
                "parent_shape_id": None,
            }],
            "messages": [],
        })

    monkeypatch.setattr(office_adapter, "_run_powershell_json", fake_probe)

    context = office_adapter.OfficeAdapter().read_context({
        "class_name": "PPTFrameClass",
        "title": "live.pptx - PowerPoint",
        "hwnd": 4242,
    })

    assert "ActivePresentation" not in captured["script"]
    assert "4242" in captured["script"]
    assert context.method == "com:powerpoint.selection"
    assert context.content == "Unsaved slide text"
    assert context.artifacts["document_saved"] is False
    assert context.artifacts["locators"][0] == {
        "kind": "slide-shape",
        "value": {
            "slideId": 512,
            "shapeId": 7,
            "parentShapeId": None,
            "bboxPt": [72.0, 144.0, 216.0, 72.0],
        },
    }


def test_excel_native_reader_binds_hwnd_and_exposes_cell_locator(monkeypatch) -> None:
    captured: dict[str, str] = {}

    def fake_probe(script: str, *, timeout: int = 2) -> office_adapter.OfficeProbeResult:
        captured["script"] = script
        return office_adapter.OfficeProbeResult(True, {
            "app": "excel",
            "method": "com:excel.selection",
            "hwnd": 8181,
            "workbook": "D:\\Sheets\\live.xlsx",
            "workbook_saved": False,
            "worksheet": "Forecast",
            "address": "B3:C4",
            "rows": [[{"text": "125", "value": 125, "formula": ""}]],
            "row_count": 2,
            "col_count": 2,
            "messages": [],
        })

    monkeypatch.setattr(office_adapter, "_run_powershell_json", fake_probe)
    context = office_adapter.OfficeAdapter().read_context({
        "class_name": "XLMAIN",
        "title": "live.xlsx - Excel",
        "hwnd": 8181,
    })

    assert "$targetHwnd = [int64]8181" in captured["script"]
    assert "candidate.HWND" in captured["script"]
    assert context.artifacts["document_saved"] is False
    assert context.artifacts["locators"] == [{
        "kind": "cell-range",
        "value": {"workbook": "D:\\Sheets\\live.xlsx", "sheet": "Forecast", "range": "B3:C4"},
    }]


def test_word_native_reader_exposes_revision_and_text_locator(monkeypatch) -> None:
    monkeypatch.setattr(
        office_adapter,
        "_run_word_selection_vbs",
        lambda prog_id, timeout=3: office_adapter.OfficeProbeResult(True, {
            "hwnd": 9191,
            "document": "D:\\Docs\\live.docx",
            "document_name": "live.docx",
            "document_path": "D:\\Docs",
            "document_saved": False,
            "selection_type": "2",
            "selection_start": 40,
            "selection_end": 58,
            "text": "Unsaved Word text",
        }),
    )

    context = office_adapter.OfficeAdapter().read_context({
        "class_name": "OpusApp",
        "title": "live.docx - Word",
        "hwnd": 9191,
    })

    assert context.artifacts["document_saved"] is False
    assert context.artifacts["source_identity"] == {
        "absolutePath": "D:\\Docs\\live.docx",
        "hwnd": 9191,
        "host": "microsoft_word",
    }
    assert context.artifacts["locators"] == [{
        "kind": "text",
        "value": {"story": "selection", "start": 40, "end": 58},
    }]


def test_word_rejects_active_document_from_another_window(monkeypatch) -> None:
    monkeypatch.setattr(
        office_adapter,
        "_run_word_selection_vbs",
        lambda prog_id, timeout=3: office_adapter.OfficeProbeResult(True, {
            "hwnd": 1111,
            "document": "D:\\Docs\\wrong.docx",
            "text": "wrong window",
        }),
    )
    monkeypatch.setattr(
        office_adapter,
        "_run_powershell_json",
        lambda script, timeout=2: office_adapter.OfficeProbeResult(True, {
            "hwnd": 2222,
            "document": "D:\\Docs\\right.docx",
            "document_name": "right.docx",
            "document_path": "D:\\Docs",
            "document_saved": True,
            "selection_start": 0,
            "selection_end": 5,
            "text": "right",
            "messages": [],
        }),
    )

    context = office_adapter.OfficeAdapter().read_context({
        "class_name": "OpusApp",
        "title": "right.docx - Word",
        "hwnd": 2222,
    })

    assert context.content == "right"
    assert context.artifacts["hwnd"] == 2222
