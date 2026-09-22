from pathlib import Path
import fitz
import pytest
from docx import Document
from pptx import Presentation
from app.actions.office_document import OfficeDocumentActionHandler, PowerShellLiveOfficeGateway
from app.actions.document_output import DocumentOutputHandler, _verify_pptx
from app.actions.pdf import PdfActionHandler
from app.actions.figma import FigmaActionHandler
from app.artifacts.document_patch import PatchOperation
from app.context_pack.sources import SourceRef


def source(path: Path, **identity):
    return SourceRef(source_id="s", task_id="t", kind="document", title=path.name,
                     identity={"absolutePath": str(path), **identity}, revision={},
                     capabilities=("read", "patch"), origin="user-attached", parent_source_id=None)


def op(name, kind, locator, before, after, identity="op"):
    return PatchOperation.from_dict({"operationId": identity, "operation": name,
        "referenceId": "ref", "sourceId": "s", "locator": {"kind": kind, "value": locator},
        "before": before, "after": after})


def test_word_replacement_at_run_boundary_keeps_target_format(tmp_path):
    path = tmp_path / "mixed.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Keep ").bold = True
    p.add_run("old").italic = True
    p.add_run(" ending").underline = True
    doc.save(path)
    operation = op("replace_text", "text", {"paragraphIndex": 0}, "Keep old ending", "Keep new ending")
    assert OfficeDocumentActionHandler().execute(source(path), operation).ok
    runs = Document(path).paragraphs[0].runs
    assert runs[0].text == "Keep "
    assert runs[1].text == "new" and runs[1].italic
    assert runs[2].text == " ending" and runs[2].underline


def test_live_word_writes_minimal_range(monkeypatch):
    calls = []
    monkeypatch.setattr(PowerShellLiveOfficeGateway, "_run", staticmethod(lambda host, body, payload: calls.append((body, payload)) or {"ok": True}))
    PowerShellLiveOfficeGateway().replace_word(path="file.docx", hwnd=1, start=20, end=35,
        expected_text="Keep old ending", replacement="Keep new ending")
    body, payload = calls[0]
    assert payload["changeStart"] == 25
    assert payload["changeEnd"] == 28
    assert payload["replacement"] == "new"
    assert "$document.Range([int]$p.changeStart, [int]$p.changeEnd)" in body


@pytest.mark.parametrize("after", [[[1]], [[1, 2], [3]], [[1, 2], [3, 4], [5, 6]]])
def test_live_excel_rejects_shape_before_native_write(monkeypatch, after):
    calls = []
    monkeypatch.setattr(PowerShellLiveOfficeGateway, "_run", staticmethod(lambda *args: calls.append(args) or {"ok": True}))
    with pytest.raises(ValueError, match="shape"):
        PowerShellLiveOfficeGateway().set_excel(path="file.xlsx", hwnd=1, sheet="Sheet1", address="A1:B2", expected=[[0, 0], [0, 0]], after=after)
    assert not calls


def annotation(path, identity, rect):
    before = {"annotationId": identity, "kind": "text-note", "outputPath": str(path), "present": False, "text": "wanted"}
    return op("add_pdf_annotation", "pdf-region", {"pageIndex": 0, "rectPt": rect}, before, {**before, "present": True}, identity)


def test_pdf_batch_continues_bound_output_but_rejects_unrelated_file(tmp_path):
    original, output = tmp_path / "original.pdf", tmp_path / "annotated.pdf"
    with fitz.open() as doc:
        doc.new_page()
        doc.save(original)
    first, second = annotation(output, "first", [40, 40, 90, 60]), annotation(output, "second", [120, 40, 170, 60])
    handler = PdfActionHandler()
    assert handler.execute(source(original), first).ok
    assert PdfActionHandler().read_current(source(original), second).value == second.before
    assert PdfActionHandler().execute(source(original), second).ok
    with fitz.open(output) as doc:
        assert len(list(doc[0].annots())) == 2
    unrelated = tmp_path / "unrelated.pdf"
    unrelated.write_bytes(original.read_bytes())
    assert not handler.execute(source(original), annotation(unrelated, "x", [1, 1, 20, 20])).ok


def test_pdf_readback_observes_content_and_geometry(tmp_path):
    original, output = tmp_path / "original.pdf", tmp_path / "annotated.pdf"
    with fitz.open() as doc:
        doc.new_page()
        doc.save(original)
    operation = annotation(output, "one", [40, 40, 90, 60])
    handler = PdfActionHandler()
    assert handler.execute(source(original), operation).ok
    assert handler.read_current(source(original), operation).value == operation.after
    with fitz.open(output) as doc:
        page = doc[0]
        annot = next(page.annots())
        annot.set_info(content="changed outside MP")
        annot.update()
        doc.saveIncr()
    assert handler.read_current(source(original), operation).value["text"] == "changed outside MP"
    with fitz.open(output) as doc:
        page = doc[0]
        annot = next(page.annots())
        annot.set_rect(fitz.Rect(200, 200, 218, 218))
        annot.update()
        doc.saveIncr()
    assert handler.read_current(source(original), operation).value != operation.after


class Figma:
    task_id = "t"
    document_session_id = "d"
    def __init__(self, text): self.text = text
    def request(self, method, payload):
        if method == "apply_patch":
            operation = payload["operations"][0]
            assert self.text[operation["start"]:operation["end"]] == operation["before"]
            self.text = self.text[:operation["start"]] + operation["after"] + self.text[operation["end"]:]
            return {"ok": True}
        return {"nodes": [{"id": "n", "characters": self.text}]}


@pytest.mark.parametrize("after", ["LONGER", "x", ""])
def test_figma_reads_changed_range_length(tmp_path, after):
    client = Figma("abcDEFghi")
    operation = op("replace_text", "figma-node", {"nodeId": "n", "textStart": 3, "textEnd": 6}, "DEF", after)
    handler = FigmaActionHandler(client)
    assert handler.execute(source(tmp_path / "x", documentSessionId="d"), operation).ok
    assert client.text == "abc" + after + "ghi"
    assert handler.read_current(source(tmp_path / "x", documentSessionId="d"), operation).value == after


def test_figma_missing_end_means_remaining_text(tmp_path):
    client = Figma("abcdef")
    operation = op("replace_text", "figma-node", {"nodeId": "n", "textStart": 3}, "def", "new suffix")
    assert FigmaActionHandler(client).execute(source(tmp_path / "x", documentSessionId="d"), operation).ok
    assert client.text == "abcnew suffix"


@pytest.mark.parametrize("names", [[None], ["long" * 10], ["same", "same"], ["_Sources"]])
def test_xlsx_normalizes_names_consistently(tmp_path, names):
    path = tmp_path / "new.xlsx"
    operation = op("create_file", "text", {"path": str(path)}, {"exists": False, "path": str(path)},
        {"exists": True, "path": str(path), "format": "xlsx", "references": [{"sourceId": "s"}],
         "content": {"sheets": [{"name": name, "rows": [[i]]} for i, name in enumerate(names)]}})
    result = DocumentOutputHandler().execute(source(path), operation)
    assert result.ok, result.error
    assert DocumentOutputHandler().read_current(source(path), operation).value == operation.after


def test_generated_ppt_paginate_tables_and_check_their_contents(tmp_path):
    path = tmp_path / "new.pptx"
    content = {"slides": [{"title": "Tables", "body": "Body", "tables": [{"rows": [[f"table {i}"]]} for i in range(3)]}]}
    operation = op("create_file", "text", {"path": str(path)}, {"exists": False, "path": str(path)},
        {"exists": True, "path": str(path), "format": "pptx", "references": [], "content": content})
    result = DocumentOutputHandler().execute(source(path), operation)
    assert result.ok, result.error
    doc = Presentation(path)
    tables = [shape for slide in doc.slides for shape in slide.shapes if shape.has_table]
    assert len(tables) == 3
    assert all(shape.top + shape.height <= doc.slide_height and shape.left + shape.width <= doc.slide_width for shape in tables)
    tables[1].table.cell(0, 0).text = "wrong"
    doc.save(path)
    assert not _verify_pptx(path, content, [])
