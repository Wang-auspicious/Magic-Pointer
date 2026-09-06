"""Self-contained dependency smoke used by development and packaged Python."""

from __future__ import annotations

import json
import tempfile
from pathlib import Path

import fitz
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches


def verify_document_dependencies() -> dict[str, str]:
    """Create and reopen one minimal file with every document dependency."""
    with tempfile.TemporaryDirectory(prefix="mp-document-smoke-") as raw_directory:
        directory = Path(raw_directory)

        pdf_path = directory / "probe.pdf"
        pdf = fitz.open()
        pdf.new_page().insert_text((72, 72), "runtime pdf target")
        pdf.save(pdf_path)
        pdf.close()
        reopened_pdf = fitz.open(pdf_path)
        try:
            pdf_text = reopened_pdf[0].get_text().strip()
        finally:
            reopened_pdf.close()

        docx_path = directory / "probe.docx"
        document = Document()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "runtime docx table target"
        document.save(docx_path)
        docx_text = Document(docx_path).tables[0].cell(0, 0).text.strip()

        pptx_path = directory / "probe.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        shape.text = "runtime pptx target"
        presentation.save(pptx_path)
        reopened_presentation = Presentation(pptx_path)
        pptx_text = next(
            item.text.strip()
            for item in reopened_presentation.slides[0].shapes
            if getattr(item, "has_text_frame", False) and item.text.strip()
        )

        xlsx_path = directory / "probe.xlsx"
        workbook = Workbook()
        workbook.active["B2"] = "runtime xlsx target"
        workbook.save(xlsx_path)
        reopened_workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
        try:
            xlsx_text = str(reopened_workbook.active["B2"].value or "").strip()
        finally:
            reopened_workbook.close()

    result = {
        "pdf": pdf_text,
        "docx": docx_text,
        "pptx": pptx_text,
        "xlsx": xlsx_text,
    }
    expected = {
        "pdf": "runtime pdf target",
        "docx": "runtime docx table target",
        "pptx": "runtime pptx target",
        "xlsx": "runtime xlsx target",
    }
    if result != expected:
        raise RuntimeError(f"document dependency smoke mismatch: {result!r}")
    return result


def main() -> int:
    print(json.dumps(verify_document_dependencies(), ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
