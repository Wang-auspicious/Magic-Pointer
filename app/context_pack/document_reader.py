"""Task-scoped, structure-preserving readers for local documents.

The reader deliberately builds a small transient index per call.  It does not
watch the filesystem or create a machine-wide knowledge base: a SourceRef is
the authorization and its revision remains visible in every returned fragment.
"""

from __future__ import annotations

import json
import mimetypes
import re
import time
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation

from .sources import Coverage, FragmentLocator, ReadFragment, ReadResult, SourceRef

_TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".py", ".js", ".ts", ".tsx",
    ".jsx", ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".csv",
    ".tsv", ".log", ".bat", ".ps1", ".html", ".htm", ".css", ".xml",
    ".svg",
}
_OFFICE_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx"}
_UNIT_RE = re.compile(r"(?:（|\()\s*([^（）()]{1,12}?)\s*(?:）|\))")
_QUERY_TERM_RE = re.compile(r"\w+", re.UNICODE)


@dataclass(frozen=True, slots=True)
class _Unit:
    locator: FragmentLocator
    text: str
    metadata: dict[str, Any]


@dataclass(frozen=True, slots=True)
class _ParsedDocument:
    units: tuple[_Unit, ...]
    backend: str
    structure: dict[str, Any]
    logical_units: int
    missing_reason: str | None = None
    degraded: bool = False


def _path_for(source: SourceRef) -> Path:
    raw = source.identity.get("absolutePath") or source.identity.get("path")
    if not isinstance(raw, str) or not raw.strip():
        raise ValueError(f"source {source.source_id} has no local absolutePath")
    return Path(raw).resolve()


def _cursor_offset(cursor: str | None, prefix: str) -> int:
    if cursor is None:
        return 0
    expected = f"{prefix}:"
    if not cursor.startswith(expected):
        raise ValueError(f"invalid {prefix} cursor: {cursor}")
    try:
        value = int(cursor[len(expected):])
    except ValueError as exc:
        raise ValueError(f"invalid {prefix} cursor: {cursor}") from exc
    if value < 0:
        raise ValueError(f"invalid {prefix} cursor: {cursor}")
    return value


def _decode_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8-replace"


def _paragraph_text(paragraph: Any) -> str:
    # python-docx exposes hyperlink text in recent releases, while collecting
    # w:t nodes also keeps it for older supported builds.
    return "".join(node.text or "" for node in paragraph._p.xpath(".//w:t")).strip()


def _pt(value: Any) -> float:
    return round(float(value or 0) / 12700.0, 3)


def _metadata_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return " ".join(f"{key} {_metadata_text(item)}" for key, item in value.items())
    if isinstance(value, (list, tuple)):
        return " ".join(_metadata_text(item) for item in value)
    return str(value)


class DocumentReader:
    """Read and search an explicitly registered local source.

    ``ocr_page`` is optional and page-scoped.  It receives ``(path,
    zero_based_page_index, png_bytes)`` and returns recognized text.  Without
    it, image-only PDF pages are reported as a coverage limitation instead of
    being mistaken for empty pages.
    """

    def __init__(
        self,
        *,
        max_fragment_chars: int = 16_000,
        max_result_chars: int = 48_000,
        ocr_page: Callable[[Path, int, bytes], str] | None = None,
    ) -> None:
        self.max_fragment_chars = max(256, int(max_fragment_chars))
        self.max_result_chars = max(4_096, int(max_result_chars))
        self.ocr_page = ocr_page

    def _parse(self, source: SourceRef) -> _ParsedDocument:
        path = _path_for(source)
        if not path.exists():
            raise FileNotFoundError(path)
        if path.is_dir():
            return self._parse_directory(path)
        suffix = path.suffix.casefold()
        if suffix == ".pdf":
            return self._parse_pdf(path)
        if suffix == ".docx":
            return self._parse_docx(path, source)
        if suffix == ".pptx":
            return self._parse_pptx(path)
        if suffix == ".xlsx":
            return self._parse_xlsx(path)
        mime, _ = mimetypes.guess_type(str(path))
        if suffix in _TEXT_EXTENSIONS or str(mime or "").startswith("text/"):
            return self._parse_text(path)
        return _ParsedDocument(
            units=(),
            backend="document.unsupported",
            structure={"kind": "unsupported", "suffix": suffix},
            logical_units=0,
            missing_reason=f"unsupported-file-type:{suffix or mime or 'unknown'}",
        )

    def _parse_pdf(self, path: Path) -> _ParsedDocument:
        document = fitz.open(path)
        units: list[_Unit] = []
        scanned_pages: list[int] = []
        used_ocr = False
        try:
            toc = document.get_toc(simple=True)
            page_structures: list[dict[str, Any]] = []
            for page_index, page in enumerate(document):
                page_meta = {
                    "pageIndex": page_index,
                    "sizePt": [round(page.rect.width, 3), round(page.rect.height, 3)],
                    "rotation": int(page.rotation),
                }
                page_structures.append(page_meta)
                blocks = page.get_text("blocks", sort=True)
                text_blocks = [block for block in blocks if len(block) >= 5 and str(block[4]).strip()]
                if not text_blocks:
                    scanned_pages.append(page_index)
                    if self.ocr_page is not None:
                        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                        recognized = str(self.ocr_page(path, page_index, pixmap.tobytes("png")) or "").strip()
                        if recognized:
                            used_ocr = True
                            text_blocks = [(0.0, 0.0, page.rect.width, page.rect.height, recognized)]
                for block_index, block in enumerate(text_blocks):
                    rect = [round(float(item), 3) for item in block[:4]]
                    locator = FragmentLocator("pdf-region", {
                        "pageIndex": page_index,
                        "rectPt": rect,
                        "blockIndex": block_index,
                    })
                    units.append(_Unit(
                        locator,
                        str(block[4]).strip(),
                        {
                            **page_meta,
                            "blockIndex": block_index,
                            "ocr": bool(not blocks and text_blocks),
                        },
                    ))
            reason = None
            if scanned_pages:
                suffix = ",".join(str(index) for index in scanned_pages)
                reason = (
                    f"ocr-text-has-no-precise-word-or-table-cells:pages={suffix}"
                    if used_ocr
                    else f"scanned-pages-require-ocr:pages={suffix}"
                )
            return _ParsedDocument(
                units=tuple(units),
                backend="document.pdf.pymupdf" + ("+ocr" if used_ocr else ""),
                structure={
                    "kind": "pdf",
                    "pageCount": len(document),
                    "pages": page_structures,
                    "toc": [
                        {"level": int(level), "title": str(title), "pageNumber": int(page)}
                        for level, title, page, *_ in toc
                    ],
                },
                logical_units=len(document),
                missing_reason=reason,
                degraded=bool(scanned_pages),
            )
        finally:
            document.close()

    def _parse_docx(self, path: Path, source: SourceRef) -> _ParsedDocument:
        document = Document(path)
        units: list[_Unit] = []
        table_index = 0
        paragraph_index = 0
        body_index = 0
        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                paragraph = Paragraph(child, document)
                text = _paragraph_text(paragraph)
                if text:
                    style = str(paragraph.style.name or "") if paragraph.style else ""
                    units.append(_Unit(
                        FragmentLocator("text", {
                            "story": "body",
                            "bodyIndex": body_index,
                            "paragraphIndex": paragraph_index,
                            "revision": dict(source.revision),
                        }),
                        text,
                        {
                            "story": "body",
                            "bodyIndex": body_index,
                            "paragraphIndex": paragraph_index,
                            "style": style,
                            "headingLevel": int(style.split()[-1]) if style.startswith("Heading ") and style.split()[-1].isdigit() else None,
                        },
                    ))
                paragraph_index += 1
                body_index += 1
            elif tag == "tbl":
                table = Table(child, document)
                for row_index, row in enumerate(table.rows):
                    cells = [" ".join(cell.text.split()) for cell in row.cells]
                    units.append(_Unit(
                        FragmentLocator("table", {
                            "story": "body",
                            "bodyIndex": body_index,
                            "tableIndex": table_index,
                            "rowIndex": row_index,
                            "revision": dict(source.revision),
                        }),
                        "\t".join(cells),
                        {
                            "story": "body",
                            "bodyIndex": body_index,
                            "tableIndex": table_index,
                            "rowIndex": row_index,
                            "cells": cells,
                        },
                    ))
                table_index += 1
                body_index += 1

        for section_index, section in enumerate(document.sections):
            for story, container in (("header", section.header), ("footer", section.footer)):
                for story_index, paragraph in enumerate(container.paragraphs):
                    text = _paragraph_text(paragraph)
                    if text:
                        units.append(_Unit(
                            FragmentLocator("text", {
                                "story": story,
                                "sectionIndex": section_index,
                                "paragraphIndex": story_index,
                                "revision": dict(source.revision),
                            }),
                            text,
                            {
                                "story": story,
                                "sectionIndex": section_index,
                                "paragraphIndex": story_index,
                            },
                        ))
        return _ParsedDocument(
            units=tuple(units),
            backend="document.docx.python-docx",
            structure={
                "kind": "docx",
                "bodyItems": body_index,
                "paragraphs": paragraph_index,
                "tables": table_index,
                "sections": len(document.sections),
            },
            logical_units=body_index,
            missing_reason="floating-objects-and-revisions-not-expanded",
        )

    def _parse_pptx(self, path: Path) -> _ParsedDocument:
        presentation = Presentation(path)
        units: list[_Unit] = []

        def visit_shape(shape: Any, *, slide_id: int, slide_index: int, parent_shape_id: int | None) -> None:
            shape_id = int(shape.shape_id)
            locator_value: dict[str, Any] = {
                "slideId": slide_id,
                "shapeId": shape_id,
                "parentShapeId": parent_shape_id,
                "shapeName": str(shape.name or ""),
                "bboxPt": [_pt(shape.left), _pt(shape.top), _pt(shape.width), _pt(shape.height)],
            }
            base_metadata = {
                **locator_value,
                "slideIndex": slide_index,
                "shapeType": int(shape.shape_type),
            }
            if getattr(shape, "has_text_frame", False):
                paragraphs: list[str] = []
                runs: list[dict[str, Any]] = []
                for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
                    paragraph_text = "".join(run.text for run in paragraph.runs) or paragraph.text
                    if paragraph_text.strip():
                        paragraphs.append(paragraph_text.strip())
                    for run_index, run in enumerate(paragraph.runs):
                        if run.text:
                            runs.append({
                                "paragraphIndex": paragraph_index,
                                "runIndex": run_index,
                                "text": run.text,
                                "bold": run.font.bold,
                                "italic": run.font.italic,
                                "fontName": run.font.name,
                                "fontSizePt": float(run.font.size.pt) if run.font.size else None,
                            })
                text = "\n".join(paragraphs).strip()
                if text:
                    units.append(_Unit(
                        FragmentLocator("slide-shape", dict(locator_value)),
                        text,
                        {**base_metadata, "contentKind": "text", "runs": runs},
                    ))
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows):
                    cells = [" ".join(cell.text.split()) for cell in row.cells]
                    table_locator = {**locator_value, "tableRowIndex": row_index}
                    units.append(_Unit(
                        FragmentLocator("slide-shape", table_locator),
                        "\t".join(cells),
                        {**base_metadata, "contentKind": "table", "tableRowIndex": row_index, "cells": cells},
                    ))
            if hasattr(shape, "shapes"):
                for child in shape.shapes:
                    visit_shape(child, slide_id=slide_id, slide_index=slide_index, parent_shape_id=shape_id)
            elif not getattr(shape, "has_text_frame", False) and not getattr(shape, "has_table", False):
                # Image/graphic names and stable identity remain discoverable;
                # their pixels are still interpreted by the visual evidence path.
                units.append(_Unit(
                    FragmentLocator("slide-shape", dict(locator_value)),
                    f"[shape] {shape.name}",
                    {**base_metadata, "contentKind": "image-or-graphic"},
                ))

        for slide_index, slide in enumerate(presentation.slides):
            slide_id = int(slide.slide_id)
            for shape in slide.shapes:
                visit_shape(shape, slide_id=slide_id, slide_index=slide_index, parent_shape_id=None)
            if getattr(slide, "has_notes_slide", False):
                notes = slide.notes_slide.notes_text_frame
                text = "\n".join(paragraph.text for paragraph in notes.paragraphs if paragraph.text.strip()).strip()
                if text:
                    units.append(_Unit(
                        FragmentLocator("text", {"story": "notes", "slideId": slide_id}),
                        text,
                        {"story": "notes", "slideId": slide_id, "slideIndex": slide_index},
                    ))
        return _ParsedDocument(
            units=tuple(units),
            backend="document.pptx.python-pptx",
            structure={
                "kind": "pptx",
                "slideCount": len(presentation.slides),
                "slideIds": [int(slide.slide_id) for slide in presentation.slides],
                "slideSizePt": [_pt(presentation.slide_width), _pt(presentation.slide_height)],
            },
            logical_units=len(presentation.slides),
        )

    @staticmethod
    def _merged_value(sheet: Any, row: int, column: int) -> Any:
        for merged in sheet.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
                return sheet.cell(merged.min_row, merged.min_col).value
        return sheet.cell(row, column).value

    def _parse_xlsx(self, path: Path) -> _ParsedDocument:
        formulas = load_workbook(path, data_only=False, read_only=False)
        cached = load_workbook(path, data_only=True, read_only=False)
        units: list[_Unit] = []
        try:
            for sheet in formulas.worksheets:
                cached_sheet = cached[sheet.title]
                data_start_row = next((
                    row_index
                    for row_index in range(1, sheet.max_row + 1)
                    if any(
                        isinstance(sheet.cell(row_index, column_index).value, (int, float))
                        or (
                            isinstance(sheet.cell(row_index, column_index).value, str)
                            and str(sheet.cell(row_index, column_index).value).startswith("=")
                        )
                        for column_index in range(1, sheet.max_column + 1)
                    )
                ), sheet.max_row + 1)
                for row in sheet.iter_rows():
                    for cell in row:
                        value = cell.value
                        if value is None:
                            continue
                        formula = value if isinstance(value, str) and value.startswith("=") else None
                        cached_value = cached_sheet[cell.coordinate].value if formula else value
                        headers: list[str] = []
                        for header_row in range(1, min(cell.row, data_start_row)):
                            candidate = self._merged_value(sheet, header_row, cell.column)
                            if candidate is not None and str(candidate).strip():
                                rendered = str(candidate).strip()
                                if rendered not in headers:
                                    headers.append(rendered)
                        row_headers: list[str] = []
                        for header_column in range(1, cell.column):
                            candidate = self._merged_value(sheet, cell.row, header_column)
                            if candidate is not None and str(candidate).strip():
                                rendered = str(candidate).strip()
                                if rendered not in row_headers:
                                    row_headers.append(rendered)
                        unit = None
                        for header in (*headers, *row_headers):
                            match = _UNIT_RE.search(header)
                            if match:
                                unit = match.group(1).strip()
                                break
                        merged_range = next((
                            str(item) for item in sheet.merged_cells.ranges
                            if item.min_row <= cell.row <= item.max_row and item.min_col <= cell.column <= item.max_col
                        ), None)
                        metadata = {
                            "sheet": sheet.title,
                            "address": cell.coordinate,
                            "headers": headers,
                            "rowHeaders": row_headers,
                            "unit": unit,
                            "formula": formula,
                            "cachedValue": cached_value if formula else value,
                            "cachedValueKnown": not formula or cached_value is not None,
                            "mergedRange": merged_range,
                            "sheetHidden": sheet.sheet_state != "visible",
                            "rowHidden": bool(sheet.row_dimensions[cell.row].hidden),
                            "columnHidden": bool(sheet.column_dimensions[get_column_letter(cell.column)].hidden),
                        }
                        rendered = [
                            f"{sheet.title}!{cell.coordinate}",
                            f"headers={' > '.join(headers)}" if headers else "",
                            f"rowHeaders={' > '.join(row_headers)}" if row_headers else "",
                            f"unit={unit}" if unit else "",
                            f"value={value}" if not formula else f"formula={formula}",
                            (
                                f"cachedValue={cached_value}"
                                if formula and cached_value is not None
                                else ("cachedValue=unknown" if formula else "")
                            ),
                        ]
                        units.append(_Unit(
                            FragmentLocator("cell-range", {"sheet": sheet.title, "range": cell.coordinate}),
                            " | ".join(item for item in rendered if item),
                            metadata,
                        ))
            return _ParsedDocument(
                units=tuple(units),
                backend="document.xlsx.openpyxl",
                structure={
                    "kind": "xlsx",
                    "sheetCount": len(formulas.worksheets),
                    "sheets": [
                        {
                            "name": sheet.title,
                            "state": sheet.sheet_state,
                            "maxRow": sheet.max_row,
                            "maxColumn": sheet.max_column,
                            "mergedRanges": [str(item) for item in sheet.merged_cells.ranges],
                        }
                        for sheet in formulas.worksheets
                    ],
                },
                logical_units=len(formulas.worksheets),
            )
        finally:
            formulas.close()
            cached.close()

    def _parse_text(self, path: Path) -> _ParsedDocument:
        text, encoding = _decode_text(path)
        lines = text.splitlines() or ([text] if text else [])
        units = tuple(
            _Unit(
                FragmentLocator("text", {"lineStart": index + 1, "lineEnd": index + 1}),
                line,
                {"lineStart": index + 1, "lineEnd": index + 1, "encoding": encoding},
            )
            for index, line in enumerate(lines)
        )
        return _ParsedDocument(
            units=units,
            backend=f"document.text.{encoding}",
            structure={"kind": "text", "lineCount": len(lines), "encoding": encoding},
            logical_units=len(lines),
        )

    def _parse_directory(self, path: Path) -> _ParsedDocument:
        entries = sorted(path.iterdir(), key=lambda item: (not item.is_dir(), item.name.casefold()))
        units: list[_Unit] = []
        for index, child in enumerate(entries):
            try:
                stat = child.stat()
                size = None if child.is_dir() else int(stat.st_size)
                mtime_ns = int(stat.st_mtime_ns)
            except OSError:
                size = None
                mtime_ns = None
            relative = child.relative_to(path).as_posix()
            units.append(_Unit(
                FragmentLocator("text", {"entryIndex": index, "relativePath": relative}),
                ("[directory] " if child.is_dir() else "[file] ") + relative,
                {
                    "entryIndex": index,
                    "relativePath": relative,
                    "absolutePath": str(child.resolve()),
                    "isDirectory": child.is_dir(),
                    "size": size,
                    "mtimeNs": mtime_ns,
                },
            ))
        return _ParsedDocument(
            units=tuple(units),
            backend="document.directory",
            structure={"kind": "directory", "entryCount": len(entries)},
            logical_units=len(entries),
        )

    def _directory_content_matches(
        self,
        source: SourceRef,
        parsed: _ParsedDocument,
        needle: str,
        terms: tuple[str, ...],
    ) -> list[tuple[bool, int, int, _Unit]]:
        """Search names first, then the readable content of each direct child.

        A directory result keeps the directory entry index as its fragment id,
        so ``follow`` can turn the hit into an independently readable SourceRef.
        The precise locator inside the child remains attached to the hit.
        """
        matches: list[tuple[bool, int, int, _Unit]] = []
        for index, entry in enumerate(parsed.units):
            entry_haystack = f"{entry.text} {_metadata_text(entry.metadata)}".casefold()
            entry_exact, entry_score = self._match_quality(entry_haystack, needle, terms)
            if entry_exact:
                matches.append((True, entry_score, index, entry))
                continue

            absolute_path = entry.metadata.get("absolutePath")
            if entry.metadata.get("isDirectory") or not isinstance(absolute_path, str):
                if entry_score:
                    matches.append((False, entry_score, index, entry))
                continue
            child_path = Path(absolute_path)
            child_source = self._child_source(source, index, child_path)
            try:
                child = self._parse(child_source)
            except (OSError, ValueError):
                continue
            child_candidates = []
            for child_index, unit in enumerate(child.units):
                exact, score = self._match_quality(
                    f"{unit.text} {_metadata_text(unit.metadata)}".casefold(),
                    needle,
                    terms,
                )
                if score:
                    child_candidates.append((exact, score, child_index, unit))
            child_match = max(
                child_candidates,
                key=lambda item: (item[0], item[1], -item[2]),
                default=None,
            )
            if child_match is None:
                if entry_score:
                    matches.append((False, entry_score, index, entry))
                continue
            child_exact, child_score, _child_index, child_unit = child_match
            if entry_score >= child_score and not child_exact:
                matches.append((False, entry_score, index, entry))
                continue
            child_locator = child_unit.locator.to_dict()
            matches.append((child_exact, child_score, index, _Unit(
                locator=FragmentLocator("text", {
                    "entryIndex": index,
                    "relativePath": entry.metadata["relativePath"],
                    "matchedChildLocator": child_locator,
                }),
                text=f"{entry.text}\n{child_unit.text}",
                metadata={
                    **entry.metadata,
                    "matchedBackend": child.backend,
                    "matchedChildLocator": child_locator,
                    "matchedChildMetadata": child_unit.metadata,
                },
            )))
        return matches

    @staticmethod
    def _match_quality(
        haystack: str,
        needle: str,
        terms: tuple[str, ...],
    ) -> tuple[bool, int]:
        if needle in haystack:
            return True, 10_000 + DocumentReader._term_score(haystack, terms)
        if len(terms) <= 1:
            return False, 0
        return False, DocumentReader._term_score(haystack, terms)

    @staticmethod
    def _term_score(haystack: str, terms: tuple[str, ...]) -> int:
        matched = tuple(term for term in terms if term in haystack)
        return len(matched) * 100 + sum(haystack.count(term) for term in matched)

    @staticmethod
    def _search_group_key(index: int, unit: _Unit) -> tuple[str, Any]:
        """Group structure that users perceive as one searchable unit."""
        if unit.locator.kind == "pdf-region":
            return "pdf-page", unit.locator.value.get("pageIndex")
        if unit.locator.kind == "slide-shape":
            return "ppt-slide", unit.locator.value.get("slideId")
        return "unit", index

    @staticmethod
    def _locator_matches(requested: FragmentLocator, actual: FragmentLocator) -> bool:
        if requested.kind != actual.kind:
            return False
        if requested.kind == "pdf-region":
            page_index = requested.value.get("pageIndex")
            if not isinstance(page_index, int):
                return False
            if actual.value.get("pageIndex") != page_index:
                return False
            block_index = requested.value.get("blockIndex")
            return block_index is None or actual.value.get("blockIndex") == block_index
        if requested.kind == "slide-shape":
            slide_id = requested.value.get("slideId")
            if slide_id is None or actual.value.get("slideId") != slide_id:
                return False
            shape_id = requested.value.get("shapeId")
            return shape_id is None or actual.value.get("shapeId") == shape_id
        return actual.value == requested.value

    def _child_source(self, source: SourceRef, index: int, path: Path) -> SourceRef:
        suffix = path.suffix.casefold()
        capabilities = ["read"]
        if path.is_dir() or suffix in (_OFFICE_EXTENSIONS | _TEXT_EXTENSIONS):
            capabilities.extend(("search", "follow"))
        if path.is_dir() or suffix in _OFFICE_EXTENSIONS:
            capabilities.append("patch")
        return SourceRef(
            source_id=f"{source.source_id}:entry:{index}",
            task_id=source.task_id,
            kind="document" if suffix in _OFFICE_EXTENSIONS else "file",
            title=path.name,
            identity={"absolutePath": str(path.resolve())},
            revision={"mtimeNs": path.stat().st_mtime_ns if path.exists() else None},
            capabilities=tuple(capabilities),
            origin="task-discovered",
            parent_source_id=source.source_id,
        )

    def _fragment(self, source: SourceRef, unit: _Unit, index: int) -> ReadFragment:
        text = unit.text
        if len(text) > self.max_fragment_chars:
            text = text[: self.max_fragment_chars].rstrip() + "\n[TRUNCATED]"
        return ReadFragment(
            fragment_id=f"fragment:{source.source_id}:unit:{index}",
            locator=unit.locator,
            text=text,
            metadata={
                **unit.metadata,
                "sourceTitle": source.title,
                "sourceRevision": dict(source.revision),
            },
            citations=({"sourceId": source.source_id, "locator": unit.locator.to_dict()},),
        )

    def _fit_result_budget(
        self,
        source: SourceRef,
        indexed_units: list[tuple[int, _Unit]],
    ) -> list[tuple[int, _Unit]]:
        """Keep a prefix whose complete structured result stays model-visible.

        The Runtime has a final 64k result guard.  Paging here, while unit
        indices and cursors are still known, prevents that guard from silently
        replacing the middle of a document with a truncation marker.
        """
        kept: list[tuple[int, _Unit]] = []
        used = 2_048
        for index, unit in indexed_units:
            fragment = self._fragment(source, unit, index)
            fragment_chars = len(json.dumps(
                fragment.to_dict(),
                ensure_ascii=False,
                separators=(",", ":"),
                default=str,
            ))
            range_chars = len(json.dumps(
                unit.locator.to_dict(),
                ensure_ascii=False,
                separators=(",", ":"),
                default=str,
            ))
            cost = fragment_chars + range_chars + 32
            if kept and used + cost > self.max_result_chars:
                break
            kept.append((index, unit))
            used += cost
        return kept

    def _result(
        self,
        source: SourceRef,
        parsed: _ParsedDocument,
        indexed_units: list[tuple[int, _Unit]],
        *,
        extent: str,
        complete: bool,
        next_cursor: str | None,
        started: float,
    ) -> ReadResult:
        fragments = tuple(self._fragment(source, unit, index) for index, unit in indexed_units)
        status = "unsupported" if parsed.backend == "document.unsupported" else (
            "degraded" if parsed.degraded else ("ok" if fragments else "empty_confirmed")
        )
        return ReadResult(
            source_id=source.source_id,
            fragments=fragments,
            coverage=Coverage(
                extent=extent,
                read_ranges=tuple(unit.locator.to_dict() for _, unit in indexed_units),
                total_units=parsed.logical_units,
                complete=complete and parsed.backend != "document.unsupported",
                next_cursor=next_cursor,
                missing_reason=parsed.missing_reason,
            ),
            evidence_status=status,
            used_backend=parsed.backend,
            latency_ms=(time.perf_counter() - started) * 1000.0,
        )

    def describe(self, source: SourceRef) -> ReadResult:
        started = time.perf_counter()
        try:
            parsed = self._parse(source)
        except Exception as exc:
            return self._error_result(source, exc, started)
        unit = _Unit(
            FragmentLocator("text", {"document": True}),
            f"{source.title}: {parsed.structure.get('kind', 'document')}",
            {"structure": parsed.structure},
        )
        return self._result(
            source, parsed, [(0, unit)], extent="document", complete=True,
            next_cursor=None, started=started,
        )

    def read(
        self,
        source: SourceRef,
        locator: FragmentLocator | None,
        cursor: str | None,
        limit: int,
    ) -> ReadResult:
        started = time.perf_counter()
        try:
            parsed = self._parse(source)
            bounded = max(1, min(int(limit), 1000))
            if locator is not None:
                matches = [
                    index for index, unit in enumerate(parsed.units)
                    if self._locator_matches(locator, unit.locator)
                ]
                if not matches:
                    return self._result(
                        source, parsed, [], extent="neighborhood", complete=False,
                        next_cursor=None, started=started,
                    )
                target = matches[0]
                target_unit = parsed.units[target]
                if locator.kind == "pdf-region":
                    page = locator.value.get("pageIndex")
                    page_order = (page, page + 1, page - 1) if isinstance(page, int) else (page,)
                    candidates = [
                        (index, unit)
                        for candidate_page in page_order
                        for index, unit in enumerate(parsed.units)
                        if unit.locator.value.get("pageIndex") == candidate_page
                    ]
                elif locator.kind == "slide-shape":
                    slide_id = locator.value.get("slideId")
                    candidates = [
                        (index, unit) for index, unit in enumerate(parsed.units)
                        if unit.locator.value.get("slideId") == slide_id
                    ]
                else:
                    start = max(0, target - 1)
                    candidates = list(enumerate(parsed.units[start:target + 2], start=start))
                if (target, target_unit) not in candidates:
                    candidates.insert(0, (target, target_unit))
                neighborhood_offset = _cursor_offset(cursor, "neighborhood")
                requested = candidates[neighborhood_offset:neighborhood_offset + bounded]
                selected = self._fit_result_budget(source, requested)
                neighborhood_end = neighborhood_offset + len(selected)
                next_cursor = (
                    f"neighborhood:{neighborhood_end}"
                    if neighborhood_end < len(candidates)
                    else None
                )
                return self._result(
                    source, parsed, selected, extent="neighborhood",
                    complete=next_cursor is None, next_cursor=next_cursor, started=started,
                )
            offset = _cursor_offset(cursor, "unit")
            requested_end = min(len(parsed.units), offset + bounded)
            requested = list(enumerate(parsed.units[offset:requested_end], start=offset))
            selected = self._fit_result_budget(source, requested)
            end = offset + len(selected)
            next_cursor = f"unit:{end}" if end < len(parsed.units) else None
            return self._result(
                source, parsed, selected, extent="document",
                complete=next_cursor is None, next_cursor=next_cursor, started=started,
            )
        except Exception as exc:
            return self._error_result(source, exc, started)

    def search(
        self,
        source: SourceRef,
        query: str,
        cursor: str | None,
        limit: int,
    ) -> ReadResult:
        started = time.perf_counter()
        if not str(query).strip():
            raise ValueError("query must be non-empty")
        try:
            parsed = self._parse(source)
            needle = str(query).strip().casefold()
            terms = tuple(dict.fromkeys(_QUERY_TERM_RE.findall(needle)))
            ranked = (
                self._directory_content_matches(source, parsed, needle, terms)
                if parsed.structure.get("kind") == "directory"
                else []
            )
            if parsed.structure.get("kind") != "directory":
                unit_haystacks = [
                    f"{unit.text} {_metadata_text(unit.metadata)}".casefold()
                    for unit in parsed.units
                ]
                grouped_haystacks: dict[tuple[str, Any], list[str]] = {}
                grouped_units: dict[tuple[str, Any], list[tuple[int, _Unit]]] = {}
                for index, (unit, haystack) in enumerate(zip(parsed.units, unit_haystacks)):
                    group_key = self._search_group_key(index, unit)
                    grouped_haystacks.setdefault(group_key, []).append(haystack)
                    grouped_units.setdefault(group_key, []).append((index, unit))
                joined_groups = {
                    key: " ".join(values)
                    for key, values in grouped_haystacks.items()
                }
                for index, (unit, haystack) in enumerate(zip(parsed.units, unit_haystacks)):
                    exact = needle in haystack
                    if exact:
                        score = 10_000 + self._term_score(haystack, terms)
                    elif len(terms) > 1:
                        group = joined_groups[self._search_group_key(index, unit)]
                        score = self._term_score(group, terms)
                    else:
                        score = 0
                    if score:
                        ranked.append((exact, score, index, unit))
            exact_matches = [item for item in ranked if item[0]]
            if exact_matches:
                ranked = exact_matches
            elif parsed.structure.get("kind") == "pdf":
                page_results: list[tuple[bool, int, int, _Unit]] = []
                for group_key, members in grouped_units.items():
                    if group_key[0] != "pdf-page":
                        continue
                    group_text = joined_groups[group_key]
                    score = self._term_score(group_text, terms) if len(terms) > 1 else 0
                    if not score:
                        continue
                    rects = [
                        unit.locator.value.get("rectPt")
                        for _index, unit in members
                        if isinstance(unit.locator.value.get("rectPt"), list)
                        and len(unit.locator.value["rectPt"]) == 4
                    ]
                    locator_value: dict[str, Any] = {"pageIndex": group_key[1]}
                    if rects:
                        locator_value["rectPt"] = [
                            min(float(rect[0]) for rect in rects),
                            min(float(rect[1]) for rect in rects),
                            max(float(rect[2]) for rect in rects),
                            max(float(rect[3]) for rect in rects),
                        ]
                    first_index, first_unit = members[0]
                    page_results.append((False, score, first_index, _Unit(
                        locator=FragmentLocator("pdf-region", locator_value),
                        text="\n".join(unit.text for _index, unit in members),
                        metadata={
                            **first_unit.metadata,
                            "searchAggregate": "pdf-page",
                            "blockIndices": [
                                unit.locator.value.get("blockIndex")
                                for _index, unit in members
                            ],
                        },
                    )))
                ranked = page_results
            ranked.sort(key=lambda item: (-item[1], item[2]))
            matches = [(index, unit) for _exact, _score, index, unit in ranked]
            offset = _cursor_offset(cursor, "match")
            bounded = max(1, min(int(limit), 1000))
            requested_end = min(len(matches), offset + bounded)
            requested = matches[offset:requested_end]
            selected = self._fit_result_budget(source, requested)
            end = offset + len(selected)
            next_cursor = f"match:{end}" if end < len(matches) else None
            return self._result(
                source, parsed, selected, extent="query-results",
                complete=next_cursor is None, next_cursor=next_cursor, started=started,
            )
        except Exception as exc:
            return self._error_result(source, exc, started)

    def follow(self, source: SourceRef, fragment_id: str) -> tuple[SourceRef, ...]:
        parsed = self._parse(source)
        prefix = f"fragment:{source.source_id}:unit:"
        if not fragment_id.startswith(prefix):
            return ()
        try:
            index = int(fragment_id[len(prefix):])
            unit = parsed.units[index]
        except (ValueError, IndexError):
            return ()
        path_value = unit.metadata.get("absolutePath")
        if not isinstance(path_value, str):
            return ()
        path = Path(path_value)
        return (self._child_source(source, index, path),)

    def preview(self, source: SourceRef, *, max_chars: int = 16_000) -> ReadResult:
        """Return a bounded first reading; callers retain the cursor to continue."""
        # A character budget and a unit budget are deliberately separate.  The
        # former protects prompts, the latter keeps structured locators intact.
        result = self.read(source, None, None, 100)
        if sum(len(fragment.text) for fragment in result.fragments) <= max_chars:
            return result
        kept: list[ReadFragment] = []
        used = 0
        for fragment in result.fragments:
            if used >= max_chars:
                break
            remaining = max_chars - used
            text = fragment.text[:remaining]
            kept.append(ReadFragment(
                fragment_id=fragment.fragment_id,
                locator=fragment.locator,
                text=text,
                metadata=fragment.metadata,
                citations=fragment.citations,
            ))
            used += len(text)
        next_index = len(kept)
        return ReadResult(
            source_id=result.source_id,
            fragments=tuple(kept),
            coverage=Coverage(
                extent=result.coverage.extent,
                read_ranges=tuple(fragment.locator.to_dict() for fragment in kept),
                total_units=result.coverage.total_units,
                complete=False,
                next_cursor=f"unit:{next_index}",
                missing_reason=result.coverage.missing_reason,
            ),
            evidence_status=result.evidence_status,
            used_backend=result.used_backend,
            latency_ms=result.latency_ms,
        )

    @staticmethod
    def _error_result(source: SourceRef, exc: Exception, started: float) -> ReadResult:
        return ReadResult(
            source_id=source.source_id,
            fragments=(),
            coverage=Coverage("document", (), None, False, None, f"{type(exc).__name__}:{exc}"),
            evidence_status="error",
            used_backend="document.error",
            latency_ms=(time.perf_counter() - started) * 1000.0,
        )


__all__ = ["DocumentReader"]
