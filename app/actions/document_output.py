"""Create editable Office files and verify them by reopening their structure."""

from __future__ import annotations

import os
import uuid
from collections.abc import Mapping, Sequence
from contextlib import suppress
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from app.artifacts.document_patch import (
    OperationReadResult,
    OperationWriteResult,
    PatchOperation,
)
from app.context_pack.sources import SourceRef

_FORMATS = frozenset({"docx", "xlsx", "pptx"})


def _state(value: Any, name: str) -> dict[str, Any]:
    if not isinstance(value, Mapping):
        raise ValueError(f"create_file {name} must be an object")
    return dict(value)


def _output_contract(operation: PatchOperation) -> tuple[dict[str, Any], dict[str, Any], Path, str]:
    if operation.operation != "create_file":
        raise ValueError(f"unsupported document output operation: {operation.operation}")
    before = _state(operation.before, "before")
    after = _state(operation.after, "after")
    if before.get("exists") is not False or after.get("exists") is not True:
        raise ValueError("create_file must transition exists false to true")
    before_path = Path(str(before.get("path") or "")).expanduser().resolve(strict=False)
    output = Path(str(after.get("path") or "")).expanduser().resolve(strict=False)
    if before_path != output:
        raise ValueError("create_file path cannot change between before and after")
    file_format = str(after.get("format") or "").casefold()
    if file_format not in _FORMATS or output.suffix.casefold() != f".{file_format}":
        raise ValueError("create_file format must match .docx, .xlsx, or .pptx path")
    if not isinstance(after.get("content"), Mapping):
        raise ValueError("create_file after.content must be an object")
    references = after.get("references")
    if not isinstance(references, list) or any(not isinstance(item, Mapping) for item in references):
        raise ValueError("create_file after.references must be an array of objects")
    return before, after, output, file_format


def _write_sources_docx(document: Document, references: Sequence[Mapping[str, Any]]) -> None:
    if not references:
        return
    document.add_heading("Sources", level=1)
    for reference in references:
        label = str(reference.get("label") or reference.get("sourceId") or "Source")
        source_id = str(reference.get("sourceId") or "")
        document.add_paragraph(f"{label} — {source_id}" if source_id else label)


def _create_docx(path: Path, content: Mapping[str, Any], references: Sequence[Mapping[str, Any]]) -> None:
    document = Document()
    title = str(content.get("title") or "").strip()
    if title:
        document.add_heading(title, level=0)
    paragraphs = content.get("paragraphs") or []
    if not isinstance(paragraphs, list):
        raise ValueError("docx paragraphs must be an array")
    for paragraph in paragraphs:
        document.add_paragraph(str(paragraph))
    tables = content.get("tables") or []
    if not isinstance(tables, list):
        raise ValueError("docx tables must be an array")
    for table_spec in tables:
        if not isinstance(table_spec, Mapping) or not isinstance(table_spec.get("rows"), list):
            raise ValueError("docx table requires rows")
        rows = list(table_spec["rows"])
        width = max((len(row) for row in rows if isinstance(row, list)), default=0)
        if not rows or width == 0:
            continue
        table = document.add_table(rows=len(rows), cols=width)
        for row_index, row in enumerate(rows):
            if not isinstance(row, list):
                raise ValueError("docx table row must be an array")
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = str(value)
    _write_sources_docx(document, references)
    document.save(path)


def _create_xlsx(path: Path, content: Mapping[str, Any], references: Sequence[Mapping[str, Any]]) -> None:
    workbook = Workbook()
    default = workbook.active
    sheets = content.get("sheets") or []
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("xlsx content requires at least one sheet")
    for index, sheet_spec in enumerate(sheets):
        if not isinstance(sheet_spec, Mapping):
            raise ValueError("xlsx sheet must be an object")
        name = str(sheet_spec.get("name") or f"Sheet{index + 1}")[:31]
        sheet = default if index == 0 else workbook.create_sheet()
        sheet.title = name
        rows = sheet_spec.get("rows") or []
        if not isinstance(rows, list):
            raise ValueError("xlsx sheet rows must be an array")
        for row in rows:
            if not isinstance(row, list):
                raise ValueError("xlsx row must be an array")
            sheet.append(list(row))
    if references:
        source_sheet = workbook.create_sheet("_Sources")
        source_sheet.append(["Label", "Source ID"])
        for reference in references:
            source_sheet.append([
                str(reference.get("label") or reference.get("sourceId") or "Source"),
                str(reference.get("sourceId") or ""),
            ])
    workbook.save(path)


def _create_pptx(path: Path, content: Mapping[str, Any], references: Sequence[Mapping[str, Any]]) -> None:
    presentation = Presentation()
    slides = content.get("slides") or []
    if not isinstance(slides, list) or not slides:
        raise ValueError("pptx content requires at least one slide")
    for slide_spec in slides:
        if not isinstance(slide_spec, Mapping):
            raise ValueError("pptx slide must be an object")
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        title = slide.shapes.title
        if title is not None:
            title.text = str(slide_spec.get("title") or "")
        body = str(slide_spec.get("body") or "")
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = body
        tables = slide_spec.get("tables") or []
        if not isinstance(tables, list):
            raise ValueError("pptx tables must be an array")
        top = 4.0
        for table_spec in tables:
            if not isinstance(table_spec, Mapping) or not isinstance(table_spec.get("rows"), list):
                raise ValueError("pptx table requires rows")
            rows = list(table_spec["rows"])
            columns = max((len(row) for row in rows if isinstance(row, list)), default=0)
            if not rows or not columns:
                continue
            table = slide.shapes.add_table(
                len(rows), columns, Inches(0.7), Inches(top), Inches(8.6), Inches(1.6)
            ).table
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = str(value)
            top += 1.8
        images = slide_spec.get("images") or []
        if not isinstance(images, list):
            raise ValueError("pptx images must be an array")
        for image in images:
            if not isinstance(image, Mapping):
                raise ValueError("pptx image must be an object")
            image_path = Path(str(image.get("path") or "")).expanduser().resolve()
            slide.shapes.add_picture(
                str(image_path),
                Inches(float(image.get("leftIn") or 0.7)),
                Inches(float(image.get("topIn") or 4.0)),
                width=Inches(float(image.get("widthIn") or 3.0)),
            )
    if references:
        source_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        if source_slide.shapes.title is not None:
            source_slide.shapes.title.text = "Sources"
        lines = [
            f"{reference.get('label') or reference.get('sourceId') or 'Source'} — {reference.get('sourceId') or ''}"
            for reference in references
        ]
        source_slide.placeholders[1].text = "\n".join(lines)
    presentation.save(path)


def _verify_docx(path: Path, content: Mapping[str, Any], references: Sequence[Mapping[str, Any]]) -> bool:
    document = Document(path)
    texts = [paragraph.text for paragraph in document.paragraphs]
    title = str(content.get("title") or "").strip()
    expected_paragraphs = [str(item) for item in content.get("paragraphs") or []]
    if title and (not texts or texts[0] != title):
        return False
    start = 1 if title else 0
    if texts[start:start + len(expected_paragraphs)] != expected_paragraphs:
        return False
    expected_tables = content.get("tables") or []
    if len(document.tables) < len(expected_tables):
        return False
    for table, expected in zip(document.tables, expected_tables, strict=False):
        actual_rows = [[cell.text for cell in row.cells] for row in table.rows]
        rows = [[str(value) for value in row] for row in expected.get("rows") or []]
        if any(
            actual[: len(wanted)] != wanted
            for actual, wanted in zip(actual_rows, rows, strict=False)
        ):
            return False
    return not references or "Sources" in texts


def _verify_xlsx(path: Path, content: Mapping[str, Any], references: Sequence[Mapping[str, Any]]) -> bool:
    workbook = load_workbook(path, data_only=False)
    try:
        for spec in content.get("sheets") or []:
            name = str(spec.get("name") or "")
            if name not in workbook.sheetnames:
                return False
            sheet = workbook[name]
            for row_index, row in enumerate(spec.get("rows") or [], start=1):
                actual = [sheet.cell(row_index, column_index).value for column_index in range(1, len(row) + 1)]
                if actual != list(row):
                    return False
        return not references or "_Sources" in workbook.sheetnames
    finally:
        workbook.close()


def _verify_pptx(path: Path, content: Mapping[str, Any], references: Sequence[Mapping[str, Any]]) -> bool:
    presentation = Presentation(path)
    expected = list(content.get("slides") or [])
    if len(presentation.slides) != len(expected) + (1 if references else 0):
        return False
    for slide, spec in zip(presentation.slides, expected, strict=False):
        texts = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
        title = str(spec.get("title") or "")
        body = str(spec.get("body") or "")
        if title not in texts or (body and body not in texts):
            return False
    if references:
        source_text = "\n".join(
            shape.text
            for shape in presentation.slides[-1].shapes
            if getattr(shape, "has_text_frame", False)
        )
        if "Sources" not in source_text:
            return False
    return True


class DocumentOutputHandler:
    used_backend = "document-output.python-office"

    def read_current(
        self,
        source: SourceRef,
        operation: PatchOperation,
    ) -> OperationReadResult:
        del source
        try:
            before, after, output, file_format = _output_contract(operation)
            if not output.exists():
                return OperationReadResult(True, before, self.used_backend)
            content = dict(after["content"])
            references = list(after["references"])
            verified = {
                "docx": _verify_docx,
                "xlsx": _verify_xlsx,
                "pptx": _verify_pptx,
            }[file_format](output, content, references)
            if verified:
                return OperationReadResult(True, after, self.used_backend)
            return OperationReadResult(
                True,
                {"exists": True, "path": str(output), "verification": "mismatch"},
                self.used_backend,
            )
        except Exception as exc:
            return OperationReadResult(
                False,
                used_backend=self.used_backend,
                error=f"document_output_read_failed:{type(exc).__name__}:{exc}",
            )

    def execute(
        self,
        source: SourceRef,
        operation: PatchOperation,
    ) -> OperationWriteResult:
        del source
        temporary: Path | None = None
        try:
            _, after, output, file_format = _output_contract(operation)
            if output.exists():
                return OperationWriteResult(
                    False, False, self.used_backend, "output_path_exists"
                )
            if not output.parent.is_dir():
                return OperationWriteResult(
                    False, False, self.used_backend, "output_parent_missing"
                )
            temporary = output.with_name(
                f".{output.stem}.{uuid.uuid4().hex}.tmp{output.suffix}"
            )
            creator = {
                "docx": _create_docx,
                "xlsx": _create_xlsx,
                "pptx": _create_pptx,
            }[file_format]
            creator(temporary, dict(after["content"]), list(after["references"]))
            verifier = {
                "docx": _verify_docx,
                "xlsx": _verify_xlsx,
                "pptx": _verify_pptx,
            }[file_format]
            if not verifier(temporary, dict(after["content"]), list(after["references"])):
                raise RuntimeError("reopen_verification_failed")
            os.replace(temporary, output)
            temporary = None
            return OperationWriteResult(True, True, self.used_backend)
        except Exception as exc:
            if temporary is not None:
                with suppress(OSError):
                    temporary.unlink(missing_ok=True)
            return OperationWriteResult(
                False,
                False,
                self.used_backend,
                f"document_output_write_failed:{type(exc).__name__}:{exc}",
            )


__all__ = ["DocumentOutputHandler"]
