"""Structured Word and Excel file edits with exact pre-read and reopen checks."""

from __future__ import annotations

import base64
import json
import os
import uuid
from collections.abc import Mapping
from contextlib import suppress
from pathlib import Path
from typing import Any, Protocol

from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import range_boundaries

from app.adapters.office_adapter import _run_powershell_json
from app.artifacts.document_patch import (
    OperationReadResult,
    OperationWriteResult,
    PatchOperation,
)
from app.context_pack.sources import SourceRef


class LiveOfficeGateway(Protocol):
    used_backend: str

    def read_word(
        self, *, path: str, hwnd: int, start: int, end: int
    ) -> str: ...

    def replace_word(
        self,
        *,
        path: str,
        hwnd: int,
        start: int,
        end: int,
        expected_text: str,
        replacement: str,
    ) -> Mapping[str, Any]: ...

    def read_excel(
        self, *, path: str, hwnd: int, sheet: str, address: str
    ) -> list[list[Any]]: ...

    def set_excel(
        self,
        *,
        path: str,
        hwnd: int,
        sheet: str,
        address: str,
        expected: list[list[Any]],
        after: list[list[Any]],
    ) -> Mapping[str, Any]: ...


def _payload_script(host: str, body: str, payload: Mapping[str, Any]) -> str:
    encoded = base64.b64encode(
        json.dumps(dict(payload), ensure_ascii=False).encode("utf-8")
    ).decode("ascii")
    prog_id = "Word.Application" if host == "word" else "Excel.Application"
    return r'''
$ErrorActionPreference = "Stop"
[Console]::OutputEncoding = [System.Text.UTF8Encoding]::new($false)
$json = [Text.Encoding]::UTF8.GetString([Convert]::FromBase64String("__PAYLOAD__"))
$p = $json | ConvertFrom-Json
$result = [ordered]@{ ok=$false; wrote=$false; error=$null; value=$null }
try {
  $application = [Runtime.InteropServices.Marshal]::GetActiveObject("__PROG_ID__")
__BODY__
} catch { $result.error = [string]$_.Exception.Message }
$result | ConvertTo-Json -Depth 16 -Compress
'''.replace("__PAYLOAD__", encoded).replace("__PROG_ID__", prog_id).replace("__BODY__", body)


class PowerShellLiveOfficeGateway:
    used_backend = "office.com.powershell"

    _WORD_BIND = r'''
  $window = $null
  foreach ($candidate in @($application.Windows)) {
    if ([int64]$candidate.Hwnd -eq [int64]$p.hwnd -and [string]::Equals(
      [string]$candidate.Document.FullName, [string]$p.path,
      [StringComparison]::OrdinalIgnoreCase
    )) { $window = $candidate; break }
  }
  if ($null -eq $window) { throw "bound_word_document_not_found" }
  $document = $window.Document
'''

    _EXCEL_BIND = r'''
  $window = $null
  $workbook = $null
  foreach ($book in @($application.Workbooks)) {
    if (-not [string]::Equals([string]$book.FullName, [string]$p.path, [StringComparison]::OrdinalIgnoreCase)) { continue }
    foreach ($candidate in @($book.Windows)) {
      if ([int64]$candidate.HWND -eq [int64]$p.hwnd) { $window = $candidate; $workbook = $book; break }
    }
    if ($null -ne $window) { break }
  }
  if ($null -eq $window) { throw "bound_excel_workbook_not_found" }
  $worksheet = $workbook.Worksheets.Item([string]$p.sheet)
  $range = $worksheet.Range([string]$p.address)
'''

    @staticmethod
    def _run(host: str, body: str, payload: Mapping[str, Any]) -> dict[str, Any]:
        probe = _run_powershell_json(_payload_script(host, body, payload), timeout=12)
        if not probe.ok:
            raise RuntimeError(probe.error or f"{host}_com_failed")
        return dict(probe.data)

    def read_word(self, *, path: str, hwnd: int, start: int, end: int) -> str:
        data = self._run("word", self._WORD_BIND + r'''
  $range = $document.Range([int]$p.start, [int]$p.end)
  $result.value = [string]$range.Text
  $result.ok = $true
''', {"path": path, "hwnd": hwnd, "start": start, "end": end})
        if data.get("ok") is not True:
            raise RuntimeError(str(data.get("error") or "word_read_failed"))
        return str(data.get("value") or "")

    def replace_word(
        self,
        *,
        path: str,
        hwnd: int,
        start: int,
        end: int,
        expected_text: str,
        replacement: str,
    ) -> Mapping[str, Any]:
        return self._run("word", self._WORD_BIND + r'''
  $range = $document.Range([int]$p.start, [int]$p.end)
  if (-not [string]::Equals([string]$range.Text, [string]$p.expected, [StringComparison]::Ordinal)) { throw "base_mismatch" }
  $range.Text = [string]$p.after
  $result.ok = $true; $result.wrote = $true
''', {
            "path": path,
            "hwnd": hwnd,
            "start": start,
            "end": end,
            "expected": expected_text,
            "after": replacement,
        })

    def read_excel(
        self, *, path: str, hwnd: int, sheet: str, address: str
    ) -> list[list[Any]]:
        data = self._run("excel", self._EXCEL_BIND + r'''
  $rows = @()
  for ($r=1; $r -le [int]$range.Rows.Count; $r++) {
    $row = @()
    for ($c=1; $c -le [int]$range.Columns.Count; $c++) {
      $cell = $range.Cells.Item($r,$c)
      $formula = [string]$cell.Formula
      $row += $(if ($formula.StartsWith("=")) { $formula } else { $cell.Value2 })
    }
    $rows += ,$row
  }
  $result.value = $rows; $result.ok = $true
''', {"path": path, "hwnd": hwnd, "sheet": sheet, "address": address})
        if data.get("ok") is not True or not isinstance(data.get("value"), list):
            raise RuntimeError(str(data.get("error") or "excel_read_failed"))
        return [list(row) if isinstance(row, list) else [row] for row in data["value"]]

    def set_excel(
        self,
        *,
        path: str,
        hwnd: int,
        sheet: str,
        address: str,
        expected: list[list[Any]],
        after: list[list[Any]],
    ) -> Mapping[str, Any]:
        return self._run("excel", self._EXCEL_BIND + r'''
  if ([int]$range.Rows.Count -ne [int]$p.expected.Count) { throw "base_mismatch" }
  for ($r=1; $r -le [int]$range.Rows.Count; $r++) {
    if ([int]$range.Columns.Count -ne [int]$p.expected[$r-1].Count) { throw "base_mismatch" }
    for ($c=1; $c -le [int]$range.Columns.Count; $c++) {
      $cell = $range.Cells.Item($r,$c)
      $formula = [string]$cell.Formula
      $current = $(if ($formula.StartsWith("=")) { $formula } else { $cell.Value2 })
      if ([string]$current -ne [string]$p.expected[$r-1][$c-1]) { throw "base_mismatch" }
    }
  }
  for ($r=1; $r -le [int]$range.Rows.Count; $r++) {
    for ($c=1; $c -le [int]$range.Columns.Count; $c++) {
      $next = $p.after[$r-1][$c-1]
      if ($next -is [string] -and $next.StartsWith("=")) { $range.Cells.Item($r,$c).Formula = $next }
      else { $range.Cells.Item($r,$c).Value2 = $next }
    }
  }
  $result.ok = $true; $result.wrote = $true
''', {
            "path": path,
            "hwnd": hwnd,
            "sheet": sheet,
            "address": address,
            "expected": expected,
            "after": after,
        })


def _source_path(source: SourceRef) -> Path:
    identity = dict(source.identity)
    raw = (
        identity.get("absolutePath")
        or identity.get("path")
        or identity.get("document")
    )
    if not str(raw or "").strip():
        raise ValueError("Office document source has no absolute path")
    return Path(str(raw)).expanduser().resolve(strict=False)


def _live_hwnd(source: SourceRef) -> int | None:
    raw = source.identity.get("hwnd")
    if raw is None or isinstance(raw, bool):
        return None
    try:
        hwnd = int(raw)
    except (TypeError, ValueError):
        return None
    return hwnd if hwnd > 0 else None


def _word_range(operation: PatchOperation) -> tuple[int, int]:
    if operation.locator.kind != "text":
        raise ValueError("live Word replacement requires a text locator")
    value = operation.locator.value
    try:
        start = int(value["start"])
        end = int(value["end"])
    except (KeyError, TypeError, ValueError) as exc:
        raise ValueError("live Word locator requires start and end") from exc
    if start < 0 or end <= start:
        raise ValueError("live Word range is invalid")
    return start, end


def _excel_locator(operation: PatchOperation) -> tuple[str, str]:
    if operation.locator.kind != "cell-range":
        raise ValueError("Excel set_cell_values requires a cell-range locator")
    value = operation.locator.value
    sheet = str(value.get("sheet") or value.get("worksheet") or "").strip()
    address = str(value.get("range") or value.get("address") or "").strip()
    if not sheet or not address:
        raise ValueError("Excel locator requires sheet and range")
    return sheet, address


def _text_state(value: Any) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, Mapping) and isinstance(value.get("text"), str):
        return str(value["text"])
    raise ValueError("replace_text before/after require text")


def _wrap_text(operation: PatchOperation, text: str) -> Any:
    return text if isinstance(operation.before, str) else {"text": text}


def _minimal_change(before: str, after: str) -> tuple[int, int, str]:
    start = 0
    maximum = min(len(before), len(after))
    while start < maximum and before[start] == after[start]:
        start += 1
    suffix = 0
    while (
        suffix < len(before) - start
        and suffix < len(after) - start
        and before[-suffix - 1] == after[-suffix - 1]
    ):
        suffix += 1
    old_length = len(before) - start - suffix
    end = len(after) - suffix if suffix else len(after)
    return start, old_length, after[start:end]


def _replace_paragraph_range(
    paragraph: Paragraph,
    *,
    start: int,
    length: int,
    replacement: str,
) -> None:
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run(replacement)
        return
    end = start + length
    cursor = 0
    start_index = 0
    start_offset = 0
    end_index = len(runs) - 1
    end_offset = len(runs[-1].text)
    found_start = False
    for index, run in enumerate(runs):
        next_cursor = cursor + len(run.text)
        if not found_start and start <= next_cursor:
            start_index = index
            start_offset = max(0, start - cursor)
            found_start = True
        if end <= next_cursor:
            end_index = index
            end_offset = max(0, end - cursor)
            break
        cursor = next_cursor
    first = runs[start_index]
    last = runs[end_index]
    prefix = first.text[:start_offset]
    suffix = last.text[end_offset:]
    if start_index == end_index:
        first.text = prefix + replacement + suffix
        return
    first.text = prefix + replacement
    for run in runs[start_index + 1:end_index]:
        run.text = ""
    last.text = suffix


def _word_paragraph(document: Document, operation: PatchOperation) -> Paragraph:
    value = operation.locator.value
    if operation.locator.kind == "text":
        try:
            return document.paragraphs[int(value["paragraphIndex"])]
        except (KeyError, IndexError, TypeError, ValueError) as exc:
            raise ValueError("Word text locator requires an existing paragraphIndex") from exc
    if operation.locator.kind == "table":
        try:
            table = document.tables[int(value["tableIndex"])]
            cell = table.cell(int(value["rowIndex"]), int(value["columnIndex"]))
            paragraph_index = int(value.get("paragraphIndex") or 0)
            return cell.paragraphs[paragraph_index]
        except (KeyError, IndexError, TypeError, ValueError) as exc:
            raise ValueError("Word table locator requires table/row/column indices") from exc
    raise ValueError("Word replace_text requires a text or table locator")


def _excel_values(path: Path, operation: PatchOperation) -> list[list[Any]]:
    if operation.locator.kind != "cell-range":
        raise ValueError("Excel set_cell_values requires a cell-range locator")
    value = operation.locator.value
    sheet_name = str(value.get("sheet") or value.get("worksheet") or "")
    address = str(value.get("range") or value.get("address") or "")
    if not sheet_name or not address:
        raise ValueError("Excel locator requires sheet and range")
    workbook = load_workbook(path, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError("Excel worksheet no longer exists")
        sheet = workbook[sheet_name]
        min_column, min_row, max_column, max_row = range_boundaries(address)
        return [
            [sheet.cell(row, column).value for column in range(min_column, max_column + 1)]
            for row in range(min_row, max_row + 1)
        ]
    finally:
        workbook.close()


def _write_excel(path: Path, temporary: Path, operation: PatchOperation) -> None:
    value = operation.locator.value
    sheet_name = str(value.get("sheet") or value.get("worksheet") or "")
    address = str(value.get("range") or value.get("address") or "")
    rows = operation.after
    if not isinstance(rows, list) or any(not isinstance(row, list) for row in rows):
        raise ValueError("set_cell_values after must be a two-dimensional array")
    workbook = load_workbook(path, data_only=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError("Excel worksheet no longer exists")
        sheet = workbook[sheet_name]
        min_column, min_row, max_column, max_row = range_boundaries(address)
        expected_height = max_row - min_row + 1
        expected_width = max_column - min_column + 1
        if len(rows) != expected_height or any(len(row) != expected_width for row in rows):
            raise ValueError("set_cell_values dimensions do not match locator range")
        for row_offset, row in enumerate(rows):
            for column_offset, cell_value in enumerate(row):
                cell = sheet.cell(min_row + row_offset, min_column + column_offset)
                if isinstance(cell, MergedCell):
                    raise ValueError("cannot write a non-anchor merged cell")
                cell.value = cell_value
        workbook.save(temporary)
    finally:
        workbook.close()


class OfficeDocumentActionHandler:
    used_backend = "office-document.python-docx+openpyxl"

    def __init__(self, *, gateway: LiveOfficeGateway | None = None) -> None:
        self.gateway = gateway or PowerShellLiveOfficeGateway()

    def read_current(
        self,
        source: SourceRef,
        operation: PatchOperation,
    ) -> OperationReadResult:
        try:
            path = _source_path(source)
            hwnd = _live_hwnd(source)
            if operation.operation == "replace_text" and path.suffix.casefold() == ".docx" and hwnd is not None:
                start, end = _word_range(operation)
                current = self.gateway.read_word(
                    path=str(path), hwnd=hwnd, start=start, end=end
                )
                value = _wrap_text(operation, current)
                if value != operation.before:
                    after_text = _text_state(operation.after)
                    current = self.gateway.read_word(
                        path=str(path),
                        hwnd=hwnd,
                        start=start,
                        end=start + len(after_text),
                    )
                    value = _wrap_text(operation, current)
            elif operation.operation == "set_cell_values" and path.suffix.casefold() == ".xlsx" and hwnd is not None:
                sheet, address = _excel_locator(operation)
                value = self.gateway.read_excel(
                    path=str(path), hwnd=hwnd, sheet=sheet, address=address
                )
            elif operation.operation == "replace_text" and path.suffix.casefold() == ".docx":
                document = Document(path)
                value = _wrap_text(operation, _word_paragraph(document, operation).text)
            elif operation.operation == "set_cell_values" and path.suffix.casefold() == ".xlsx":
                value = _excel_values(path, operation)
            else:
                raise ValueError(
                    f"unsupported Office file operation: {operation.operation} on {path.suffix}"
                )
            backend = self.gateway.used_backend if hwnd is not None else self.used_backend
            return OperationReadResult(True, value, backend)
        except Exception as exc:
            return OperationReadResult(
                False,
                used_backend=self.used_backend,
                error=f"office_document_read_failed:{type(exc).__name__}:{exc}",
            )

    def execute(
        self,
        source: SourceRef,
        operation: PatchOperation,
    ) -> OperationWriteResult:
        temporary: Path | None = None
        try:
            path = _source_path(source)
            hwnd = _live_hwnd(source)
            if hwnd is None and not path.is_file():
                raise ValueError("Office source file is missing")
            current = self.read_current(source, operation)
            if not current.ok:
                return OperationWriteResult(
                    False, False, self.used_backend, current.error
                )
            if current.value != operation.before:
                return OperationWriteResult(
                    False, False, self.used_backend, "base_mismatch"
                )
            if operation.operation == "replace_text" and path.suffix.casefold() == ".docx" and hwnd is not None:
                start, end = _word_range(operation)
                result = self.gateway.replace_word(
                    path=str(path),
                    hwnd=hwnd,
                    start=start,
                    end=end,
                    expected_text=_text_state(operation.before),
                    replacement=_text_state(operation.after),
                )
                return OperationWriteResult(
                    bool(result.get("ok")),
                    bool(result.get("wrote")),
                    self.gateway.used_backend,
                    str(result.get("error") or "") or None,
                )
            if operation.operation == "set_cell_values" and path.suffix.casefold() == ".xlsx" and hwnd is not None:
                sheet, address = _excel_locator(operation)
                if not isinstance(operation.before, list) or not isinstance(operation.after, list):
                    raise ValueError("set_cell_values before/after must be arrays")
                result = self.gateway.set_excel(
                    path=str(path),
                    hwnd=hwnd,
                    sheet=sheet,
                    address=address,
                    expected=operation.before,
                    after=operation.after,
                )
                return OperationWriteResult(
                    bool(result.get("ok")),
                    bool(result.get("wrote")),
                    self.gateway.used_backend,
                    str(result.get("error") or "") or None,
                )
            temporary = path.with_name(
                f".{path.stem}.{uuid.uuid4().hex}.tmp{path.suffix}"
            )
            if operation.operation == "replace_text" and path.suffix.casefold() == ".docx":
                document = Document(path)
                paragraph = _word_paragraph(document, operation)
                before = _text_state(operation.before)
                after = _text_state(operation.after)
                start, length, replacement = _minimal_change(before, after)
                _replace_paragraph_range(
                    paragraph,
                    start=start,
                    length=length,
                    replacement=replacement,
                )
                document.save(temporary)
            elif operation.operation == "set_cell_values" and path.suffix.casefold() == ".xlsx":
                _write_excel(path, temporary, operation)
            else:
                raise ValueError(
                    f"unsupported Office file operation: {operation.operation} on {path.suffix}"
                )
            os.replace(temporary, path)
            temporary = None
            verified = self.read_current(source, operation)
            if not verified.ok or verified.value != operation.after:
                return OperationWriteResult(
                    False, True, self.used_backend, "write_readback_mismatch"
                )
            return OperationWriteResult(True, True, self.used_backend)
        except Exception as exc:
            if temporary is not None:
                with suppress(OSError):
                    temporary.unlink(missing_ok=True)
            return OperationWriteResult(
                False,
                False,
                self.used_backend,
                f"office_document_write_failed:{type(exc).__name__}:{exc}",
            )


__all__ = [
    "LiveOfficeGateway",
    "OfficeDocumentActionHandler",
    "PowerShellLiveOfficeGateway",
]
